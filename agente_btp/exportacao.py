"""Relatório Excel — direto e autoexplicativo, derivado de UMA estrutura (os ResultadoWO auditados).

Abas:
  Dados    — TODAS as WOs consultadas: WO, Material SAP, Tipo Material, Fornecedor SAP, Quantidade faltante,
             Mensagem e Observação ("-" quando está tudo certo; em português claro quando não).
             Célula sem informação leva "-" (nunca fica em branco). O mesmo item que o BTP
             devolve mais de uma vez com dados idênticos aparece numa linha só, com a observação.
             Tipo Material vem da tabela de valores importada (Configurações).
  Materiais — um material SAP por linha: Tipo Material, quantidade total faltante (soma de todas
             as WOs), valor total (quantidade total × unitário) e o valor unitário da tabela.
  Análise  — números da análise, materiais mais recorrentes, mensagens e a situação de cada WO.
  Rastreio — (oculta) cada leitura feita ao BTP, para investigar uma divergência se preciso.
  Tabela de valores — (oculta) a tabela importada inteira (Material SAP com 18 dígitos, como no
             BTP; Tipo Material; valor unitário), para quem quiser usá-la em fórmulas.

Ao final, o arquivo gravado é reaberto e as abas Dados e Materiais são conferidas célula a célula.
"""
from __future__ import annotations

import math
from collections import Counter, OrderedDict
from datetime import datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .materiais import TabelaMateriais, chave_material
from .modelos import CAMPOS, ROTULO, STATUS_FALHA, Auditoria, LinhaMensagem, ResultadoWO, Status

ABAS = ["Dados", "Materiais", "Análise", "Rastreio", "Tabela de valores"]
COLUNAS_TABELA = ["Material SAP", "Tipo Material", "Valor unitário (R$)"]
COLUNAS_DADOS = ["WO", "Material SAP", "Tipo Material", "Fornecedor SAP", "Quantidade faltante", "Mensagem",
                 "Observação", "Status"]
# Ao lado da quantidade, o valor total dela (quantidade total × unitário); o unitário vem depois, para conferência.
COLUNAS_MATERIAIS = ["Material SAP", "Tipo Material", "Quantidade total", "Valor total (R$)",
                     "Valor unitário (R$)", "Observação"]
_COL_OBS = COLUNAS_DADOS.index("Observação")
_REAIS = '"R$" #,##0.00'

# Status da WO na análise (o mesmo da base de execução e da tela).
ROTULO_ESTADO = {"validada": "Validada", "divergente": "Divergente", "incompleta": "Incompleta",
                 "falha": "Falha", "pendente": "Pendente"}

VERMELHO = "DA291C"
_CAB_FONTE = Font(bold=True, color="FFFFFF")
_CAB_FUNDO = PatternFill("solid", fgColor=VERMELHO)
_ATENCAO = PatternFill("solid", fgColor="FFF2CC")       # só a célula de Observação, quando há algo
_BORDA = Border(bottom=Side(style="thin", color="D9D9D9"))
_QUEBRA = Alignment(wrap_text=True, vertical="top")
_TOPO = Alignment(vertical="top")

# Situação da WO em linguagem direta (sem jargão).
SITUACAO = {
    Status.PROCESSADA: "Com itens faltantes",
    Status.SEM_REGISTROS: "No Monitor, sem itens na janela Mensagens",
    Status.NAO_ENCONTRADA: "Nenhum alerta encontrado no Monitor para esta WO",
    Status.WO_INVALIDA: "Número de WO inválido",
    Status.SESSAO_EXPIRADA: "Não consultada: sessão do BTP expirou",
    Status.FALHA_COMUNICACAO: "Não consultada: falha de comunicação",
    Status.TEMPO_ESGOTADO: "Não consultada: tempo esgotado",
    Status.NAO_PROCESSADA: "Pendente: ainda não consultada",
    Status.ERRO: "Não consultada: erro na consulta",
}


def nome_analise(pasta: str | Path, numero: int, total_wos: int) -> Path:
    """Ex.: 'Análise 12 - 499 WOs.xlsx'. Nunca sobrescreve um arquivo existente."""
    base = f"Análise {numero} - {total_wos} WO{'s' if total_wos != 1 else ''}"
    destino = Path(pasta) / f"{base}.xlsx"
    extra = 2
    while destino.exists():
        destino = Path(pasta) / f"{base} ({extra}).xlsx"
        extra += 1
    return destino


def _txt(v) -> str:
    if v is None or v == "":
        return ""
    if isinstance(v, Decimal):
        return format(v.normalize(), "f")
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def observacao_item(l: LinhaMensagem) -> str:
    """Observação em português claro; vazia quando o item foi conferido sem problema."""
    if l.auditoria is Auditoria.VALIDADO:
        return ""
    partes = []
    for campo in CAMPOS:
        nota = l.notas.get(campo, "")
        if not nota:
            continue
        rotulo = ROTULO[campo]
        if nota.startswith("Divergente"):
            valores = nota.split(":", 1)[1].split(".")[0].strip()
            valores = " e ".join(v.split("=", 1)[-1] for v in valores.split(" / "))
            if "Resolvida" in nota:
                partes.append(f"{rotulo} mudou no BTP durante a consulta ({valores}); vale o valor atual")
            else:
                partes.append(f"{rotulo} diferente entre consultas ({valores}): conferir no BTP")
        elif nota.startswith("Ambígua"):
            partes.append(f"{rotulo}: o BTP tem mais de um campo de quantidade; conferir no BTP")
        elif nota.startswith("Não obtido"):
            partes.append(f"{rotulo} não informado pelo BTP")
        else:
            partes.append(f"{rotulo}: {nota}")
    if "ausente" in l.observacao or "Não estava" in l.observacao:
        partes.append("Item apareceu/sumiu entre consultas: conferir no BTP")
    if l.auditoria is Auditoria.NAO_CONCLUIDO and not partes:
        partes.append("Conferência não concluída (" + (l.observacao or "tempo ou falha") + ")")
    return "; ".join(partes) + ("." if partes else "")


def estado_rotulo(r: ResultadoWO) -> str:
    if r.status is Status.NAO_PROCESSADA:
        return ROTULO_ESTADO["pendente"]
    from .loja import estado_da_wo
    return ROTULO_ESTADO[estado_da_wo(r)]


SEM_INFO = "-"          # célula sem informação: nunca em branco
NOTA_REPETIDO = "O BTP traz este item {n} vezes nesta WO, com os mesmos dados (alertas diferentes); listado uma vez."


def itens_unicos(linhas: list[LinhaMensagem]) -> list[tuple[LinhaMensagem, int]]:
    """Itens da WO sem repetição: o MESMO item (material, fornecedor, quantidade, mensagem e
    conferência idênticos) que o BTP devolve mais de uma vez vira UMA linha, com a contagem.
    Nada é somado nem escolhido: as cópias são iguais em tudo que vai para a planilha."""
    vistos: dict[tuple, list] = OrderedDict()
    for l in linhas:
        chave = (l.wo, l.material, l.fornecedor, _txt(l.quantidade_faltante), l.mensagem, observacao_item(l))
        if chave in vistos:
            vistos[chave][1] += 1
        else:
            vistos[chave] = [l, 1]
    return [(l, n) for l, n in vistos.values()]


def _preencher(valores: list) -> list:
    return [SEM_INFO if v is None or v == "" else v for v in valores]


def _tipo(tabela: TabelaMateriais | None, material) -> str | None:
    achado = tabela.buscar(material) if tabela else None
    return achado.tipo if achado else None


def linhas_dados(resultados: list[ResultadoWO], exportar_incompletos: bool = True,
                 tabela: TabelaMateriais | None = None) -> list[list]:
    """Linhas da aba Dados: um item por linha; WO sem itens aparece numa linha com a observação."""
    saida = []
    for r in resultados:
        estado = estado_rotulo(r)
        itens = [l for l in r.linhas if exportar_incompletos or l.auditoria is Auditoria.VALIDADO]
        for l, n in itens_unicos(itens):
            obs = "; ".join(o for o in (observacao_item(l), NOTA_REPETIDO.format(n=n) if n > 1 else "") if o)
            saida.append(_preencher([l.wo, l.material, _tipo(tabela, l.material), l.fornecedor,
                                     l.quantidade_faltante, l.mensagem, obs, estado]))
        if not itens:
            obs = SITUACAO[r.status]
            if r.linhas:   # havia itens, mas não validados e a exportação deles está desligada
                obs = f"{len(r.linhas)} item(ns) com divergência omitido(s): conferir no BTP"
            elif r.status is Status.NAO_PROCESSADA:
                obs = r.observacao or obs
            elif r.status in STATUS_FALHA and r.observacao and r.observacao != "OK":
                obs = f"{obs} ({r.observacao})"
            saida.append(_preencher([r.wo, None, None, None, None, None, obs, estado]))
    return saida


def _decimal(v) -> Decimal:
    return v if isinstance(v, Decimal) else Decimal(str(v))


def _numero(d: Decimal):
    """Decimal exato para gravar: inteiro quando não tem casas (12, não 12.0)."""
    return int(d) if d == d.to_integral_value() else d.normalize()


def linhas_materiais(itens: list[LinhaMensagem], tabela: TabelaMateriais | None) -> tuple[list[list], list]:
    """Aba Materiais: um material SAP por linha (todas as WOs juntas) e a linha de total.

    Quantidade total = soma da quantidade faltante dos itens CONFERIDOS (os mesmos da aba Dados,
    sem as cópias). Item não conferido ou sem quantidade fica fora da soma e é avisado.
    Valor total = quantidade total × valor unitário da tabela. Sem valor na tabela: "-"."""
    grupos: dict[str, list[LinhaMensagem]] = OrderedDict()
    for l in itens:
        if l.material:
            grupos.setdefault(chave_material(l.material), []).append(l)
    linhas = []
    for ls in grupos.values():
        somados = [l for l in ls if l.auditoria is Auditoria.VALIDADO and l.quantidade_faltante is not None]
        quantidade = sum((_decimal(l.quantidade_faltante) for l in somados), Decimal(0)) if somados else None
        achado = tabela.buscar(ls[0].material) if tabela else None
        unitario = achado.unitario if achado else None
        total = quantidade * unitario if quantidade is not None and unitario is not None else None
        obs = []
        if tabela is None:
            obs.append("Tabela de valores não importada (Configurações)")
        elif achado is None:
            obs.append("Material fora da tabela de valores")
        elif unitario is None:
            obs.append("Sem valor unitário na tabela de valores")
        fora = len(ls) - len(somados)
        if fora:
            obs.append(f"{fora} item(ns) não conferido(s) fora da soma: conferir no BTP")
        linhas.append(_preencher([ls[0].material, achado.tipo if achado else None,
                                  None if quantidade is None else _numero(quantidade),
                                  None if total is None else _numero(total),
                                  None if unitario is None else _numero(unitario), "; ".join(obs) + ("." if obs else "")]))
    linhas.sort(key=lambda l: (l[3] == SEM_INFO, -(l[3] if l[3] != SEM_INFO else 0), l[0]))
    sem_valor = sum(1 for l in linhas if l[3] == SEM_INFO)
    soma = sum((_decimal(l[3]) for l in linhas if l[3] != SEM_INFO), Decimal(0))
    total = _preencher(["Total", None, None, _numero(soma) if len(linhas) > sem_valor else None, None,
                        f"{sem_valor} material(is) sem valor ficaram fora do total." if sem_valor else None])
    return linhas, total


def gerar_planilha(resultados: list[ResultadoWO], destino: str | Path, numero: int | None = None,
                   emitida_em: datetime | None = None, origem: str = "", prefs=None,
                   fonte_campos: dict | None = None, releitura_por_chave: bool | None = None,
                   tabela: TabelaMateriais | None = None) -> Path:
    """Grava o Excel em modo de ESCRITA CONTÍNUA (write-only): cada linha vai direto para o
    arquivo, sem montar a planilha inteira na memória (13 mil WOs: ~1 GB a menos)."""
    destino = Path(destino)
    destino.parent.mkdir(parents=True, exist_ok=True)
    emitida_em = emitida_em or datetime.now()
    exportar_incompletos = True if prefs is None else prefs.exportar_incompletos
    # Itens como na aba Dados (o mesmo item repetido pelo BTP conta uma vez: não dobra quantidade).
    linhas = [l for r in resultados for l, _ in itens_unicos(
        [l for l in r.linhas if exportar_incompletos or l.auditoria is Auditoria.VALIDADO])]
    wb = Workbook(write_only=True)

    # ================================================================ Dados
    ws = wb.create_sheet(ABAS[0])
    dados = linhas_dados(resultados, exportar_incompletos, tabela)
    _preparar(ws, [18, 22, 34, 16, 12, 55, 55, 13], linha_cab=1)
    if dados:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUNAS_DADOS))}{len(dados) + 1}"
    ws.append(_cabecalho(ws, COLUNAS_DADOS))
    for valores in dados:
        # Item com observação de conferência: destaca só a célula de Observação.
        obs = valores[_COL_OBS]
        destaque = valores[1] != SEM_INFO and obs != SEM_INFO and not obs.startswith("O BTP traz")
        ws.append([_celula(ws, v, formato="@" if col in (1, 2, 3, 4, 6, 7, 8) else None,
                           alinhamento=_QUEBRA if col in (6, 7) else _TOPO, borda=_BORDA,
                           fundo=_ATENCAO if col == _COL_OBS + 1 and destaque else None)
                   for col, v in enumerate(valores, start=1)])

    # ================================================================ Materiais
    wm = wb.create_sheet(ABAS[1])
    mats, total = linhas_materiais(linhas, tabela)
    _preparar(wm, [22, 44, 16, 18, 20, 55], linha_cab=1)
    if mats:
        wm.auto_filter.ref = f"A1:{get_column_letter(len(COLUNAS_MATERIAIS))}{len(mats) + 1}"
    wm.append(_cabecalho(wm, COLUNAS_MATERIAIS))
    for valores in mats + [None, total]:
        if valores is None:
            wm.append([])
            continue
        negrito = Font(bold=True) if valores is total else None
        wm.append([_celula(wm, v, fonte=negrito, borda=_BORDA, alinhamento=_QUEBRA if col == 6 else _TOPO,
                           formato="@" if col in (1, 2) else (_REAIS if col in (4, 5) and v != SEM_INFO else None),
                           fundo=_ATENCAO if col == 6 and v != SEM_INFO and valores is not total else None)
                   for col, v in enumerate(valores, start=1)])

    # ================================================================ Análise
    wa = wb.create_sheet(ABAS[2])
    for col, largura in enumerate([40, 30, 12, 40, 16, 50], start=1):
        wa.column_dimensions[get_column_letter(col)].width = largura
    wa.append([_celula(wa, f"Análise{' nº ' + str(numero) if numero else ''} — Monitor de Alertas",
                       fonte=Font(bold=True, size=14, color=VERMELHO))])
    wa.append([_celula(wa, f"Emitida em {emitida_em:%d/%m/%Y %H:%M}" + (f" · {origem}" if origem else ""),
                       fonte=Font(size=10, color="595959"))])
    wa.append([_celula(wa, (f"Tabela de valores: {tabela.arquivo} (importada em "
                            f"{datetime.fromisoformat(tabela.importada_em):%d/%m/%Y %H:%M}, {len(tabela.itens)} materiais)")
                       if tabela else "Tabela de valores: não importada (Tipo Material e valores ficam \"-\")",
                       fonte=Font(size=10, color="595959"))])
    st = Counter(r.status for r in resultados)
    com_obs = sum(1 for l in linhas if l.auditoria is not Auditoria.VALIDADO)
    _secao(wa, "Números", ["Indicador", "Quantidade"], [
        ("WOs consultadas", len(resultados)),
        ("WOs com itens faltantes", st[Status.PROCESSADA]),
        ("WOs no Monitor sem itens na janela Mensagens", st[Status.SEM_REGISTROS]),
        ("WOs não encontradas no Monitor de Alertas", st[Status.NAO_ENCONTRADA]),
        ("WOs não consultadas (erro, tempo ou sessão)", sum(st[s] for s in STATUS_FALHA if s is not Status.NAO_PROCESSADA)),
        ("WOs pendentes (análise ainda não concluída)", st[Status.NAO_PROCESSADA]),
        ("WOs com número inválido", st[Status.WO_INVALIDA]),
        ("Itens faltantes (linhas com material na aba Dados)", len(linhas)),
        ("Itens conferidos sem diferença", len(linhas) - com_obs),
        ("Itens com observação (conferir no BTP)", com_obs),
    ])

    grupos: dict[tuple, list[LinhaMensagem]] = OrderedDict()
    for l in linhas:
        grupos.setdefault((l.material or "(não informado)", l.fornecedor or "(não informado)"), []).append(l)
    materiais = []
    for (material, fornecedor), ls in sorted(grupos.items(), key=lambda kv: (-len({l.wo for l in kv[1]}), kv[0])):
        wos = list(OrderedDict.fromkeys(l.wo for l in ls))
        materiais.append((material, fornecedor, len(wos), ", ".join(wos), _soma(ls),
                          " | ".join(OrderedDict.fromkeys(l.mensagem or "(sem mensagem)" for l in ls))))
    _secao(wa, "Materiais (mais recorrentes primeiro)",
           ["Material SAP", "Fornecedor SAP", "Nº de WOs", "WOs", "Quantidade faltante total", "Mensagens"],
           materiais, texto={1, 2})

    msgs = Counter(l.mensagem or "(sem mensagem)" for l in linhas)
    wos_por_msg: dict[str, set] = {}
    for l in linhas:
        wos_por_msg.setdefault(l.mensagem or "(sem mensagem)", set()).add(l.wo)
    _secao(wa, "Mensagens", ["Mensagem", "Itens", "WOs"],
           [(m, n, len(wos_por_msg[m])) for m, n in msgs.most_common()])

    por_wo = []
    for r in resultados:
        ls = [l for l, _ in itens_unicos([l for l in r.linhas if exportar_incompletos or l.auditoria is Auditoria.VALIDADO])]
        obs = "; ".join(o for o in OrderedDict.fromkeys(observacao_item(l) for l in ls) if o)
        por_wo.append((r.wo, SITUACAO[r.status], len(ls),
                       ", ".join(OrderedDict.fromkeys(l.material for l in ls if l.material)),
                       obs or ("" if r.status not in STATUS_FALHA else r.observacao)))
    _secao(wa, "Situação de cada WO", ["WO", "Situação", "Itens", "Materiais SAP", "Observação"], por_wo, texto={1})

    # ================================================================ Rastreio (oculta)
    wr = wb.create_sheet(ABAS[3])
    wr.sheet_state = "hidden"
    wr.append(_cabecalho(wr, ["WO", "Material SAP", "Identificação do item", "Leitura", "Tipo", "Data/hora",
                              "Quantidade", "Fornecedor SAP", "Mensagem", "Comparação", "Resultado"]))
    for r in resultados:
        for a in r.registros_auditoria:
            wr.append([_celula(wr, v) for v in (
                a["wo"], a["material"], a["ocorrencia"], a["rodada"], a["tipo"],
                a["quando"].replace(microsecond=0) if a.get("quando") else None, a["quantidade"],
                a["fornecedor"], a["mensagem"], a["resultado"], a["situacao"])])
    if fonte_campos:
        wr.append([])
        wr.append([_celula(wr, v) for v in ["Campos do BTP usados"]
                   + [f"{ROTULO[c]} = {fonte_campos.get(c, '?')}" for c in CAMPOS]])

    # ================================================================ Tabela de valores (oculta)
    wt = wb.create_sheet(ABAS[4])
    wt.sheet_state = "hidden"
    _preparar(wt, [22, 44, 18], linha_cab=1)
    wt.append(_cabecalho(wt, COLUNAS_TABELA))
    for valores in linhas_tabela(tabela):
        wt.append([_celula(wt, v, formato="@" if col in (1, 2) else (_REAIS if col == 3 and v != SEM_INFO else None))
                   for col, v in enumerate(valores, start=1)])

    wb.save(destino)
    return destino


def linhas_tabela(tabela: TabelaMateriais | None) -> list[list]:
    """A tabela de valores importada, com o material no formato do BTP (18 dígitos) para que
    PROCV/XLOOKUP com a coluna Material SAP da aba Dados funcione direto."""
    if not tabela:
        return []
    return [_preencher([chave.zfill(18) if chave.isdigit() else chave, m.tipo,
                        None if m.unitario is None else _numero(m.unitario)])
            for chave, m in sorted(tabela.itens.items(), key=lambda kv: (len(kv[0]), kv[0]))]


def _soma(linhas: list[LinhaMensagem]):
    """Soma só itens conferidos sem diferença e com quantidade; senão None (não chuta)."""
    if not linhas or any(l.auditoria is not Auditoria.VALIDADO or l.quantidade_faltante is None for l in linhas):
        return None
    total = sum((Decimal(str(l.quantidade_faltante)) for l in linhas), Decimal(0))
    return int(total) if total == total.to_integral_value() else total


def _secao(ws, titulo, cabecalho, linhas, texto=frozenset()):
    """Seção da aba Análise: linha em branco, título, cabeçalho e as linhas (em sequência)."""
    ws.append([])
    ws.append([_celula(ws, titulo, fonte=Font(bold=True, size=12))])
    ws.append(_cabecalho(ws, cabecalho))
    for valores in linhas:
        ws.append([_celula(ws, v, alinhamento=_QUEBRA, borda=_BORDA, formato="@" if col in texto else None)
                   for col, v in enumerate(valores, start=1)])
    if not linhas:
        ws.append([_celula(ws, "(nenhum)", fonte=Font(italic=True, color="808080"))])


def verificar_planilha(destino: str | Path, resultados: list[ResultadoWO],
                       exportar_incompletos: bool = True, tabela: TabelaMateriais | None = None) -> list[str]:
    """Reabre o arquivo gravado e confere as abas Dados e Materiais, célula a célula."""
    esperadas = linhas_dados(resultados, exportar_incompletos, tabela)
    itens = [l for r in resultados for l, _ in itens_unicos(
        [l for l in r.linhas if exportar_incompletos or l.auditoria is Auditoria.VALIDADO])]
    mats, total = linhas_materiais(itens, tabela)
    # Somente leitura: lê célula a célula sem montar a planilha inteira na memória (13 mil WOs:
    # 3 MB em vez de ~840 MB) e ainda informa o TIPO de cada célula (fórmula = "f").
    wb = load_workbook(destino, read_only=True)
    try:
        if wb.sheetnames != ABAS:
            return [f"Abas fora da ordem esperada: {wb.sheetnames}"]
        problemas: list[str] = []
        dados = _ler_aba(wb[ABAS[0]], len(COLUNAS_DADOS), problemas)
        materiais = _ler_aba(wb[ABAS[1]], len(COLUNAS_MATERIAIS), problemas)
        tabela_gravada = _ler_aba(wb[ABAS[4]], len(COLUNAS_TABELA), problemas)
    finally:
        wb.close()
    problemas += _comparar("Dados", COLUNAS_DADOS, dados, esperadas)
    problemas += _comparar("Materiais", COLUNAS_MATERIAIS, materiais,
                           mats + ([[None] * len(COLUNAS_MATERIAIS), total]))
    problemas += _comparar("Tabela de valores", COLUNAS_TABELA, tabela_gravada, linhas_tabela(tabela))
    return problemas


def _ler_aba(ws, colunas: int, problemas: list[str]) -> list[tuple]:
    linhas = []
    for i, linha in enumerate(ws.iter_rows(max_col=colunas), start=1):
        valores = tuple(c.value for c in linha)
        linhas.append(valores + (None,) * (colunas - len(valores)))
        problemas += [f"{ws.title}!{get_column_letter(col)}{i}: gravado como fórmula ({c.value!r})"
                      for col, c in enumerate(linha, start=1) if getattr(c, "data_type", "") == "f"]
    return linhas


def _comparar(aba: str, cabecalho: list[str], linhas: list[tuple], esperadas: list[list]) -> list[str]:
    if not linhas or list(linhas[0]) != cabecalho:
        return [f"Cabeçalho da aba {aba} diferente: {linhas[0] if linhas else '(vazia)'}"]
    problemas = []
    gravadas = linhas[1:]
    if len(gravadas) != len(esperadas):
        problemas.append(f"Aba {aba} tem {len(gravadas)} linha(s); esperado {len(esperadas)}.")
    for i, (g, e) in enumerate(zip(gravadas, esperadas), start=2):
        for col, (vg, ve) in enumerate(zip(g, [_limpar(v) for v in e])):
            numeros = isinstance(vg, (int, float)) and isinstance(ve, (int, float, Decimal))
            if numeros and math.isclose(float(vg), float(ve), rel_tol=1e-12):
                continue   # o Excel guarda número com 15 dígitos significativos
            if _txt(vg) != _txt(ve):
                problemas.append(f"{aba}, linha {i}, {cabecalho[col]}: gravado {vg!r}, esperado {ve!r}")
    return problemas


def _limpar(valor):
    """Remove caracteres de controle que o formato .xlsx não aceita (derrubariam a gravação)."""
    return ILLEGAL_CHARACTERS_RE.sub("", valor) if isinstance(valor, str) else valor


def _celula(ws, valor, fonte=None, fundo=None, alinhamento=None, borda=None, formato=None) -> WriteOnlyCell:
    """Uma célula gravada SEM interpretar nada: texto que começa com "=" fica texto (nunca
    fórmula); caracteres de controle que o .xlsx não aceita são removidos."""
    c = WriteOnlyCell(ws, value=_limpar(valor))
    if isinstance(c.value, str) and c.data_type == "f":
        c.data_type = "s"
    if fonte:
        c.font = fonte
    if fundo:
        c.fill = fundo
    if alinhamento:
        c.alignment = alinhamento
    if borda:
        c.border = borda
    if formato:
        c.number_format = formato
    return c


def _cabecalho(ws, colunas: list[str]) -> list:
    saida = []
    for nome in colunas:
        c = _celula(ws, nome or None, alinhamento=Alignment(vertical="center", wrap_text=True))
        if nome:
            c.font, c.fill = _CAB_FONTE, _CAB_FUNDO
        saida.append(c)
    return saida


def _preparar(ws, larguras: list[int], linha_cab: int) -> None:
    """Larguras e painel congelado: no modo de escrita contínua, ANTES da primeira linha."""
    for i, largura in enumerate(larguras, start=1):
        ws.column_dimensions[get_column_letter(i)].width = largura
    ws.freeze_panes = f"A{linha_cab + 1}"
