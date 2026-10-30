"""Tabela de valores dos materiais (importada pelo usuário): Material SAP -> Tipo Material e
valor unitário (R$).

Vem de uma planilha do usuário (ex.: "TABELA VALORES 07-2026.xlsx", colunas Material,
Texto breve material, Unitário). Fica gravada neste computador (config/tabela_materiais.json)
e vale para toda planilha gerada depois, inclusive o relatório gerado de novo pelo Histórico.

Regras (nada é deduzido):
  - Tipo Material = "Texto breve material" da tabela, exatamente como está.
  - O material é comparado pelo NÚMERO: "000000000022056365" (BTP) = 22056365 (tabela).
  - Material fora da tabela: sem tipo e sem valor ("-" na planilha), com aviso.
  - O mesmo material duas vezes na tabela com tipo ou valor diferentes: a importação é recusada.
"""
from __future__ import annotations

import json
import os
import tempfile
import unicodedata
from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook

from . import config as cfg

COLUNA_MATERIAL = "Material"
COLUNA_TIPO = "Texto breve material"
COLUNA_UNITARIO = "Unitário"
LINHAS_PROCURA_CABECALHO = 10


def arquivo_tabela() -> Path:
    return cfg.ARQUIVO_CONFIG.parent / "tabela_materiais.json"


class TabelaInvalida(ValueError):
    """A planilha escolhida não serve como tabela de valores (motivo em português)."""


def chave_material(valor) -> str:
    """Material comparado pelo número (sem zeros à esquerda); código com letras: como está."""
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    texto = str(valor).strip()
    if texto.isdigit():
        return texto.lstrip("0") or "0"
    return texto.upper()


def _normalizar(texto) -> str:
    sem_acento = unicodedata.normalize("NFKD", str(texto or "")).encode("ascii", "ignore").decode()
    return " ".join(sem_acento.lower().split())


def _valor(v) -> Decimal | None:
    """Valor unitário como está na célula (0.13 -> 0.13, sem resíduo de ponto flutuante)."""
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, Decimal)):
        return Decimal(v)
    if isinstance(v, float):
        return Decimal(repr(v))
    texto = str(v).strip().replace("R$", "").replace(" ", "")
    if "," in texto:
        texto = texto.replace(".", "").replace(",", ".")
    try:
        return Decimal(texto) if texto else None
    except InvalidOperation:
        return None


@dataclass
class Material:
    tipo: str
    unitario: Decimal | None


@dataclass
class TabelaMateriais:
    itens: dict[str, Material]
    arquivo: str = ""
    importada_em: str = ""
    sem_valor: int = 0

    def buscar(self, material) -> Material | None:
        return self.itens.get(chave_material(material)) if material else None

    def resumo(self) -> dict:
        return {"carregada": True, "arquivo": self.arquivo, "importada_em": self.importada_em,
                "materiais": len(self.itens), "sem_valor": self.sem_valor}


def ler_planilha(caminho: str | Path, nome: str = "") -> TabelaMateriais:
    """Lê a planilha do usuário (somente leitura). Recusa com o motivo se não servir."""
    try:
        wb = load_workbook(caminho, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 — arquivo ilegível vira mensagem clara
        raise TabelaInvalida(f"Não foi possível abrir a planilha ({exc}).") from exc
    try:
        procurados = {_normalizar(COLUNA_MATERIAL): "material", _normalizar(COLUNA_TIPO): "tipo",
                      _normalizar(COLUNA_UNITARIO): "unitario"}
        for ws in wb.worksheets:
            linhas = ws.iter_rows(values_only=True)
            for _ in range(LINHAS_PROCURA_CABECALHO):
                cab = next(linhas, None)
                if cab is None:
                    break
                posicoes = {procurados[_normalizar(c)]: i for i, c in enumerate(cab) if _normalizar(c) in procurados}
                if len(posicoes) == 3:
                    return _montar(linhas, posicoes, nome or Path(caminho).name, ws.title)
    finally:
        wb.close()
    raise TabelaInvalida(f'A planilha precisa das colunas "{COLUNA_MATERIAL}", "{COLUNA_TIPO}" e '
                         f'"{COLUNA_UNITARIO}" (no cabeçalho, nas {LINHAS_PROCURA_CABECALHO} primeiras linhas).')


def _montar(linhas, pos: dict, nome: str, aba: str) -> TabelaMateriais:
    itens: dict[str, Material] = {}
    conflitos = []
    for linha in linhas:
        def celula(campo):
            return linha[pos[campo]] if pos[campo] < len(linha) else None
        chave = chave_material(celula("material"))
        if not chave:
            continue
        tipo = str(celula("tipo") or "").strip()
        novo = Material(tipo, _valor(celula("unitario")))
        atual = itens.get(chave)
        if atual and atual != novo:
            conflitos.append(chave)
        itens.setdefault(chave, novo)
    if conflitos:
        raise TabelaInvalida(f"{len(conflitos)} material(is) aparecem mais de uma vez com tipo ou valor diferentes "
                             f"(ex.: {', '.join(conflitos[:5])}). Corrija a planilha e importe de novo.")
    if not itens:
        raise TabelaInvalida(f'A aba "{aba}" tem o cabeçalho, mas nenhum material abaixo dele.')
    return TabelaMateriais(itens, nome, datetime.now().isoformat(timespec="seconds"),
                           sum(1 for m in itens.values() if m.unitario is None))


def salvar(tabela: TabelaMateriais, destino: Path | None = None) -> None:
    destino = destino or arquivo_tabela()
    destino.parent.mkdir(parents=True, exist_ok=True)
    dados = {"arquivo": tabela.arquivo, "importada_em": tabela.importada_em, "sem_valor": tabela.sem_valor,
             "itens": {k: [m.tipo, None if m.unitario is None else str(m.unitario)] for k, m in tabela.itens.items()}}
    fd, tmp = tempfile.mkstemp(dir=destino.parent, suffix=".tmp")
    with os.fdopen(fd, "w", encoding="utf-8") as f:
        json.dump(dados, f, ensure_ascii=False)
    os.replace(tmp, destino)       # gravação atômica: nunca fica uma tabela pela metade


def carregar(origem: Path | None = None) -> TabelaMateriais | None:
    """Tabela gravada, ou None se nunca foi importada. Arquivo corrompido: erro (não some calado)."""
    origem = origem or arquivo_tabela()
    if not origem.exists():
        return None
    dados = json.loads(origem.read_text(encoding="utf-8"))
    itens = {k: Material(t, None if u is None else Decimal(u)) for k, (t, u) in dados["itens"].items()}
    return TabelaMateriais(itens, dados.get("arquivo", ""), dados.get("importada_em", ""), dados.get("sem_valor", 0))


def remover(origem: Path | None = None) -> None:
    (origem or arquivo_tabela()).unlink(missing_ok=True)
