"""Leitura e validação das WOs informadas pelo usuário."""
from __future__ import annotations

import re
from pathlib import Path

# Formato observado nos prints: 5 dígitos | 9 dígitos (ex.: 02235|230771207).
PADRAO_WO = re.compile(r"^\d{5}\|\d{9}$")


def normalizar_wo(valor) -> str:
    """Converte para texto sem alterar dígitos (preserva zeros à esquerda)."""
    if valor is None:
        return ""
    texto = str(valor).strip()
    # Tolera espaços ao redor da barra: "02235 | 230771207".
    return re.sub(r"\s*\|\s*", "|", texto)


def wo_valida(wo: str) -> bool:
    return bool(PADRAO_WO.match(wo))


def ler_wos_texto(texto: str) -> list[str]:
    """Uma WO por linha; linhas vazias ignoradas; duplicadas removidas (ordem mantida)."""
    wos = [normalizar_wo(linha) for linha in texto.splitlines()]
    return _sem_duplicadas([w for w in wos if w])


def ler_wos_planilha(caminho: str | Path) -> list[str]:
    """Lê as WOs de uma planilha, tolerando variações de layout.

    1. coluna cujo cabeçalho é exatamente 'WO' (em qualquer aba);
    2. senão, coluna cujo cabeçalho contém a palavra WO (ex.: 'Nº WO', 'WO Claro');
    3. senão, qualquer célula da planilha no formato de WO (5 dígitos|9 dígitos).
    """
    from openpyxl import load_workbook

    wb = load_workbook(caminho, read_only=True, data_only=True)
    try:
        abas = [list(ws.iter_rows(values_only=True)) for ws in wb.worksheets]
    finally:
        wb.close()

    for criterio in (lambda h: h == "WO", lambda h: re.search(r"\bWO\b", h) is not None):
        for linhas in abas:
            for num, linha in enumerate(linhas[:10]):          # cabeçalho nas 10 primeiras linhas
                cab = [str(c).strip().upper() if c is not None else "" for c in linha]
                indice = next((i for i, h in enumerate(cab) if criterio(h)), None)
                if indice is not None:
                    wos = [normalizar_wo(l[indice]) for l in linhas[num + 1:] if indice < len(l)]
                    wos = [w for w in wos if w]
                    if wos:
                        return _sem_duplicadas(wos)

    encontradas = [normalizar_wo(c) for linhas in abas for linha in linhas for c in linha
                   if c is not None and PADRAO_WO.match(normalizar_wo(c))]
    if not encontradas:
        raise ValueError("Não encontrei WOs na planilha (formato esperado: 02235|230771207).")
    return _sem_duplicadas(encontradas)


def _sem_duplicadas(itens: list[str]) -> list[str]:
    vistos: set[str] = set()
    saida = []
    for item in itens:
        if item not in vistos:
            vistos.add(item)
            saida.append(item)
    return saida
