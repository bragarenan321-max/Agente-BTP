import json
from decimal import Decimal

import pytest
from openpyxl import Workbook, load_workbook

from agente_btp import config as cfg
from agente_btp.cliente import ClienteBTP, Resposta
from agente_btp.entrada import ler_wos_planilha, ler_wos_texto, wo_valida
from agente_btp.exportacao import gerar_planilha
from agente_btp.modelos import Status
from agente_btp.odata import como_quantidade, extrair_registros, preencher_modelo
from agente_btp.processador import processar_lote

CONFIG = cfg.a_partir_de_dict(json.loads(
    (cfg.RAIZ / "config" / "api.exemplo.json").read_text(encoding="utf-8")))

ALERTA = {"WO": "88412|126579742", "Consolidado": "105012-20260916-abc"}
MENSAGEM = {"WO": "88412|126579742", "Material": "000000000022026169",
            "Fornecedor": "0000137628", "QuantidadeFaltante": "2.000",
            "Mensagem": "Estoque indisponível"}


def json_v2(registros):
    return Resposta(200, "application/json", json.dumps({"d": {"results": registros}}))


class TransporteFalso:
    """Responde por trecho de URL; registra as URLs pedidas."""

    def __init__(self, rotas):
        self.rotas, self.urls = rotas, []

    def __call__(self, url):
        self.urls.append(url)
        for trecho, resposta in self.rotas.items():
            if trecho in url:
                return resposta(url) if callable(resposta) else resposta
        return json_v2([])


# ---------- entrada ----------
def test_wo_valida_preserva_zero_e_formato():
    assert wo_valida("02235|230771207")
    assert not wo_valida("2235|230771207")
    assert not wo_valida("230771207")


def test_ler_wos_texto_normaliza_e_remove_duplicadas():
    texto = " 02235 | 230771207 \n\n02235|230771207\n42340|231184065"
    assert ler_wos_texto(texto) == ["02235|230771207", "42340|231184065"]


def test_ler_wos_planilha(tmp_path):
    wb = Workbook()
    wb.active.append(["Outra", "WO"])
    wb.active.append(["x", "02235|230771207"])
    wb.active.append(["y", None])
    caminho = tmp_path / "in.xlsx"
    wb.save(caminho)
    assert ler_wos_planilha(caminho) == ["02235|230771207"]


# ---------- odata ----------
def test_extrair_registros_v2_e_v4():
    assert extrair_registros('{"d":{"results":[{"a":1}]}}') == [{"a": 1}]
    assert extrair_registros('{"value":[{"a":1}]}') == [{"a": 1}]


def test_preencher_modelo_codifica_valores():
    url = preencher_modelo("/srv/AlertasSet?$filter=WO eq '{wo}'", wo="02235|230771207")
    assert url == "/srv/AlertasSet?$filter=WO%20eq%20'02235%7C230771207'"
    url = preencher_modelo("/srv/Msg?$filter=C eq '{linha[C]}'", linha={"C": "16vp6+Z4/0CW'x"})
    assert url == "/srv/Msg?$filter=C%20eq%20'16vp6%2BZ4%2F0CW%27%27x'"


def test_paginacao_segue_next():
    paginas = {
        "AlertasSet?": Resposta(200, "application/json", json.dumps(
            {"d": {"results": [ALERTA], "__next": "AlertasSet?$skiptoken=2"}})),
        "skiptoken=2": Resposta(200, "application/json", json.dumps(
            {"d": {"results": [dict(ALERTA, Consolidado="segundo")]}})),
    }
    t = TransporteFalso({"skiptoken": paginas["skiptoken=2"], "AlertasSet": paginas["AlertasSet?"]})
    alertas = ClienteBTP(CONFIG, t).buscar_alertas("88412|126579742")
    assert [a["Consolidado"] for a in alertas] == ["105012-20260916-abc", "segundo"]


def test_quantidade_decimal_odata():
    assert como_quantidade("2.000") == 2
    assert como_quantidade("1.500") == Decimal("1.500")
    assert como_quantidade(None) is None


# ---------- estados ----------
def test_processada_com_campos_corretos():
    t = TransporteFalso({"AlertasSet": json_v2([ALERTA]), "MensagensSet": json_v2([MENSAGEM])})
    [r] = processar_lote(ClienteBTP(CONFIG, t), ["88412|126579742"])
    assert r.status is Status.PROCESSADA and r.registros == 1
    linha = r.linhas[0]
    assert linha.material == "000000000022026169"
    assert linha.fornecedor == "0000137628"
    assert linha.quantidade_faltante == 2
    assert all("$batch" not in u for u in t.urls)


def test_nao_encontrada():
    t = TransporteFalso({"AlertasSet": json_v2([])})
    [r] = processar_lote(ClienteBTP(CONFIG, t), ["88412|126579742"])
    assert r.status is Status.NAO_ENCONTRADA


def test_sem_registros():
    t = TransporteFalso({"AlertasSet": json_v2([ALERTA]), "MensagensSet": json_v2([])})
    [r] = processar_lote(ClienteBTP(CONFIG, t), ["88412|126579742"])
    assert r.status is Status.SEM_REGISTROS


def test_wo_invalida_nao_consulta():
    t = TransporteFalso({})
    [r] = processar_lote(ClienteBTP(CONFIG, t), ["123"])
    assert r.status is Status.WO_INVALIDA and t.urls == []


def test_sessao_expirada_interrompe_lote_sem_culpar_as_seguintes():
    """D1: só a WO em que a sessão caiu leva o motivo; a seguinte NÃO foi consultada (pendente)."""
    t = TransporteFalso({"AlertasSet": Resposta(200, "text/html", "<html>login</html>")})
    rs = processar_lote(ClienteBTP(CONFIG, t), ["88412|126579742", "42340|231184065"])
    assert [r.status for r in rs] == [Status.SESSAO_EXPIRADA, Status.NAO_PROCESSADA]
    assert len(t.urls) == 1


def test_redirecionamento_e_sessao_expirada():
    t = TransporteFalso({"AlertasSet": Resposta(0, "", "", redirecionado=True)})
    [r] = processar_lote(ClienteBTP(CONFIG, t), ["88412|126579742"])
    assert r.status is Status.SESSAO_EXPIRADA


def test_falha_comunicacao_continua_lote(monkeypatch):
    t = TransporteFalso({"88412": Resposta(503, "text/plain", ""),
                         "AlertasSet": json_v2([dict(ALERTA, WO="42340|231184065")]),
                         "MensagensSet": json_v2([dict(MENSAGEM, WO="42340|231184065")])})
    rs = processar_lote(ClienteBTP(CONFIG, t), ["88412|126579742", "42340|231184065"])
    assert rs[0].status is Status.FALHA_COMUNICACAO
    assert rs[1].status is Status.PROCESSADA


def test_erro_odata():
    erro = Resposta(400, "application/json", '{"error":{"message":{"value":"x"}}}')
    t = TransporteFalso({"AlertasSet": erro})
    [r] = processar_lote(ClienteBTP(CONFIG, t), ["88412|126579742"])
    assert r.status is Status.ERRO


# ---------- excel ----------
def test_planilha_preserva_zeros_e_tipos(tmp_path):
    t = TransporteFalso({"AlertasSet": json_v2([ALERTA]), "MensagensSet": json_v2([MENSAGEM])})
    rs = processar_lote(ClienteBTP(CONFIG, t), ["88412|126579742"])
    destino = gerar_planilha(rs, tmp_path / "out.xlsx", numero=7)
    wb = load_workbook(destino)
    assert wb.sheetnames == ["Dados", "Análise", "Rastreio"] and wb["Rastreio"].sheet_state == "hidden"
    ws = wb["Dados"]
    assert [c.value for c in ws[1]] == ["WO", "Material SAP", "Fornecedor SAP", "Quantidade faltante",
                                        "Mensagem", "Observação", "Status"]
    assert [c.value for c in ws[2]] == ["88412|126579742", "000000000022026169", "0000137628", 2,
                                        "Estoque indisponível", None, "Validada"]   # conferido: sem observação
    assert ws["B2"].data_type == "s" and ws["C2"].data_type == "s" and ws["D2"].data_type == "n"
    assert ws.freeze_panes == "A2" and ws.auto_filter.ref == "A1:G2"
    assert ws["E2"].alignment.wrap_text
    assert "Análise nº 7" in wb["Análise"]["A1"].value


def test_nome_da_planilha(tmp_path):
    from agente_btp.exportacao import nome_analise
    assert nome_analise(tmp_path, 12, 499).name == "Análise 12 - 499 WOs.xlsx"
    assert nome_analise(tmp_path, 3, 1).name == "Análise 3 - 1 WO.xlsx"
    (tmp_path / "Análise 3 - 1 WO.xlsx").write_bytes(b"")
    assert nome_analise(tmp_path, 3, 1).name == "Análise 3 - 1 WO (2).xlsx"


def test_config_exige_campos():
    with pytest.raises(ValueError):
        cfg.a_partir_de_dict({"url_lista": "/s/A?$filter=WO eq '{wo}'", "url_mensagens": "/s/M",
                              "campos": {"wo": "WO"}})


def test_planilha_sem_cabecalho_wo(tmp_path):
    wb = Workbook()
    wb.active.append(["Ordem de serviço", "Obs"])
    wb.active.append(["06201|739635817", "x"])
    wb.active.append([" 06282 | 328086400", None])
    caminho = tmp_path / "in.xlsx"
    wb.save(caminho)
    assert ler_wos_planilha(caminho) == ["06201|739635817", "06282|328086400"]


def test_planilha_cabecalho_com_wo_no_nome(tmp_path):
    wb = Workbook()
    wb.active.append(["Nº WO", "Outra"])
    wb.active.append(["05715|230659432", "06201|739595275"])
    caminho = tmp_path / "in.xlsx"
    wb.save(caminho)
    assert ler_wos_planilha(caminho) == ["05715|230659432"]
