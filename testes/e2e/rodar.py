"""Teste de ponta a ponta com navegador real e um SAP falso.

Rode a partir de uma CÓPIA do projeto (ele grava config/, descoberta/ e saidas/):
    python testes/e2e/rodar.py
"""
import os, sys, time
sys.path.insert(0, os.path.dirname(__file__))
os.environ["AGENTE_BTP_URL"] = "http://127.0.0.1:8798/site#Shell-home"
os.environ["AGENTE_BTP_HEADLESS"] = "1"
import sap_falso
sap_falso.iniciar(8798)
sys.path.insert(0, os.getcwd())
from agente_btp.agente import Agente, PASTA_SAIDAS
from openpyxl import load_workbook

a = Agente()
def esperar(cond, t=60):
    fim = time.time() + t
    while time.time() < fim:
        if cond(a.estado()): return True
        time.sleep(0.3)
    raise SystemExit("TIMEOUT\n" + "\n".join(a.estado()["mensagens"]))

a.enviar("conectar"); esperar(lambda e: e["conexao"] == "conectado")
a.enviar("descoberta_iniciar"); esperar(lambda e: e["descoberta_ativa"])
time.sleep(5); a.salvar_descoberta()
esperar(lambda e: not e["ocupado"] and not e["descoberta_ativa"])
a.enviar("executar", ["88412|126579742", "42340|231184065", "02235|230771207", "123"])
time.sleep(1); esperar(lambda e: not e["ocupado"] and e["ultimo_arquivo"])
print("\n".join(a.estado()["mensagens"]))
wb = load_workbook(PASTA_SAIDAS / a.estado()["ultimo_arquivo"])
for aba in wb.sheetnames:
    print(f"--- {aba}")
    for linha in wb[aba].iter_rows(values_only=True): print(linha)
print("tipo Material:", wb["Resultado"]["C2"].data_type, "| formato:", wb["Resultado"]["C2"].number_format)
a.navegador  # (thread do agente encerra com o processo)
os._exit(0)
