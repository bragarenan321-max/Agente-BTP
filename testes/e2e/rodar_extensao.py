"""Ponta a ponta da EXTENSÃO: a aba do BTP (falso) se conecta sozinha, sem favorito.

Cenários:
  1. página do BTP aberta ANTES do programa -> a extensão conecta quando o programa abre;
  2. análise completa até o Excel, sem nenhum clique na aba.

Rode a partir de uma CÓPIA do projeto:  python testes/e2e/rodar_extensao.py
"""
import json
import os
import re
import shutil
import sys
import tempfile
import threading
import time
from pathlib import Path

sys.path.insert(0, os.path.dirname(__file__))
os.environ["AGENTE_BTP_URL"] = "http://127.0.0.1:8798/site#Shell-home"
import sap_falso  # noqa: E402

sap_falso.iniciar(8798)
sys.path.insert(0, os.getcwd())
from openpyxl import load_workbook  # noqa: E402
from playwright.sync_api import sync_playwright  # noqa: E402

# Cópia da extensão apontando para o BTP falso (o resto é idêntico ao que vai para o usuário).
pasta_ext = Path(tempfile.mkdtemp()) / "extensao"
shutil.copytree("extensao", pasta_ext)
for nome in ("manifest.json", "fundo.js"):
    arq = pasta_ext / nome
    arq.write_text(re.sub(r"https://claro-prd\.launchpad\.cfapps\.br10\.hana\.ondemand\.com",
                          "http://127.0.0.1:8798", arq.read_text(encoding="utf-8")), encoding="utf-8")

with sync_playwright() as p:
    ctx = p.chromium.launch_persistent_context(
        tempfile.mkdtemp(), headless=False,
        args=["--headless=new", f"--disable-extensions-except={pasta_ext}", f"--load-extension={pasta_ext}"])
    aba = ctx.new_page()
    aba.goto("http://127.0.0.1:8798/site?csp=1#Shell-home")  # BTP aberto (sem eval), programa AINDA fechado
    aba.wait_for_timeout(3000)
    print("antes do programa, selo na aba:", aba.inner_text("#agente-btp-selo")
          if aba.locator("#agente-btp-selo").count() else "(nenhum)",
          "| aviso de erro da extensão:", aba.locator("#agente-btp-extensao-aviso").count() > 0)
    print("eval bloqueado na página:", aba.evaluate("() => { try { eval('1'); return false } catch (e) { return true } }"))

    from agente_btp.agente import PASTA_SAIDAS, Agente  # noqa: E402
    from agente_btp.servidor import criar_app  # noqa: E402
    agente = Agente()


    def _confere_amostra():  # faz o papel do usuário: confirma a amostra mostrada na tela
        while True:
            if agente.estado()["fase"] == "confirmando" and agente.estado().get("confirmacao"):
                agente.responder_fonte({"resposta": "sim", "mensagem": "item"})
            time.sleep(0.3)

    threading.Thread(target=_confere_amostra, daemon=True).start()
    app = criar_app(agente)
    threading.Thread(target=lambda: app.run(host="127.0.0.1", port=8767, threaded=True), daemon=True).start()
    inicio = time.time()

    def esperar(cond, t=60):
        fim = time.time() + t
        while time.time() < fim:
            if cond(agente.estado()):
                return
            aba.wait_for_timeout(300)
        raise SystemExit("TIMEOUT\n" + json.dumps(agente.estado(), ensure_ascii=False, indent=1))

    esperar(lambda e: e["conexao"] in ("conectado", "aba_conectada"), 30)       # nenhum clique na aba
    print(f"conectou sozinha em {time.time() - inicio:.1f} s após abrir o programa")
    print("selo na aba:", aba.inner_text("#agente-btp-selo"),
          "| aviso de erro da extensão:", aba.locator("#agente-btp-extensao-aviso").count() > 0)

    agente.enviar("executar", ["88412|126579742", "42340|231184065", "02235|230771207"])
    aba.wait_for_timeout(500)
    esperar(lambda e: e["fase"] == "concluido" and not e["ocupado"], 90)
    e = agente.estado()
    print("contagem:", e["contagem"])
    for linha in load_workbook(PASTA_SAIDAS / e["ultimo_arquivo"])["Dados"].iter_rows(values_only=True):
        print("  ", linha)
    ctx.close()
os._exit(0)
