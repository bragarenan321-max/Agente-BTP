"""Tabela de valores (Material SAP -> Tipo Material e valor unitário) e o que ela gera na planilha:
coluna Tipo Material (aba Dados), aba Materiais (soma por material) e aba oculta com a tabela."""
import io
import json
from decimal import Decimal

import pytest
from openpyxl import Workbook, load_workbook

from agente_btp import config as cfg
from agente_btp import materiais
from agente_btp.cliente import ClienteBTP, Resposta
from agente_btp.exportacao import ABAS, gerar_planilha, verificar_planilha
from agente_btp.preferencias import Preferencias
from agente_btp.processador import processar_lote

WO1, WO2, WO3, WO4 = "05659|231308367", "05456|230916326", "06187|231931856", "99999|999999999"
RAPIDO = Preferencias(intervalo_auditorias_seg=0)


def tabela_xlsx(caminho, linhas, cabecalho=("Material", "Centro Principal Parceiro", "Empresa",
                                             "Texto breve material", "UN", "Unitário", "GIRO"), antes=0):
    """Mesmo formato da tabela real (TABELA VALORES 07-2026): material numérico, R$ como número."""
    wb = Workbook()
    ws = wb.active
    for _ in range(antes):
        ws.append(["Tabela de valores"])
    ws.append(list(cabecalho))
    for l in linhas:
        ws.append(list(l))
    wb.save(caminho)
    return caminho


REAL = [(22056365, "N190", "881", "PROTETOR P/MINI ISOLADOR - ACRE1568", "UN", 2.2, "-"),
        (22056366, "N190", "881", "ISOLADOR COAXIAL QUADRADO - CISP-HR", "UN", 15.43, "-"),
        (22025349, "N190", "881", "FIXA FIO CABO DROP FO CHATO", "UN", 0.13, "-"),
        (22066302, "N190", "881", "FONTE ALIMENTACAO 1.5F18W16 120150SPAE", "UN", 36.88, "ALTO GIRO"),
        (22099999, "N190", "881", "MATERIAL SEM PRECO", "UN", None, "-")]


# ============================================================ importação
def test_le_a_tabela_no_formato_real(tmp_path):
    t = materiais.ler_planilha(tabela_xlsx(tmp_path / "t.xlsx", REAL), "TABELA VALORES 07-2026.xlsx")
    assert len(t.itens) == 5 and t.sem_valor == 1 and t.arquivo == "TABELA VALORES 07-2026.xlsx"
    m = t.buscar("000000000022056366")                     # zeros à esquerda do BTP = número da tabela
    assert m.tipo == "ISOLADOR COAXIAL QUADRADO - CISP-HR" and m.unitario == Decimal("15.43")
    assert t.buscar("000000000022025349").unitario == Decimal("0.13")   # sem resíduo de ponto flutuante
    assert t.buscar("000000000022000000") is None and t.buscar(None) is None
    assert t.buscar("000000000022099999").unitario is None


def test_cabecalho_fora_da_primeira_linha_sem_acento_e_valor_com_virgula(tmp_path):
    t = materiais.ler_planilha(tabela_xlsx(tmp_path / "u.xlsx", [(22056365, "PROTETOR", "2,20")], antes=2,
                                           cabecalho=("MATERIAL", "TEXTO BREVE MATERIAL", "Unitario")))
    assert t.buscar("22056365").tipo == "PROTETOR" and t.buscar("22056365").unitario == Decimal("2.20")


def test_recusa_tabela_sem_as_colunas(tmp_path):
    with pytest.raises(materiais.TabelaInvalida, match="Texto breve material"):
        materiais.ler_planilha(tabela_xlsx(tmp_path / "t.xlsx", [(1, "x")], cabecalho=("Material", "Descrição")))


def test_recusa_material_repetido_com_valor_diferente(tmp_path):
    linhas = REAL[:1] + [(22056365, "N190", "881", "PROTETOR P/MINI ISOLADOR - ACRE1568", "UN", 3.5, "-")]
    with pytest.raises(materiais.TabelaInvalida, match="22056365"):
        materiais.ler_planilha(tabela_xlsx(tmp_path / "t.xlsx", linhas))
    # repetido IGUAL não é conflito
    assert len(materiais.ler_planilha(tabela_xlsx(tmp_path / "u.xlsx", REAL[:1] * 2)).itens) == 1


def test_grava_e_carrega_sem_perder_precisao(tmp_path):
    t = materiais.ler_planilha(tabela_xlsx(tmp_path / "t.xlsx", REAL))
    destino = tmp_path / "config" / "tabela.json"
    materiais.salvar(t, destino)
    volta = materiais.carregar(destino)
    assert volta.itens == t.itens and volta.arquivo == t.arquivo and volta.sem_valor == 1
    assert materiais.carregar(tmp_path / "nada.json") is None
    materiais.remover(destino)
    assert not destino.exists()


# ============================================================ planilha
def _config():
    return cfg.a_partir_de_dict({
        "url_lista": "/srv/Itens?$filter=workOrderID eq '{wo}'", "campo_consolidado_lista": "",
        "url_mensagens": None, "campo_mensagens_expandidas": ".",
        "campos": {"wo": "workOrderID", "material": "material", "fornecedor": "fornecedor",
                   "quantidade_faltante": "quantidadeFaltante", "mensagem": "mensagem"}})


ITENS = {
    WO1: [("000000000022056366", 1, "i1"), ("000000000022025349", 10, "i2")],
    WO2: [("000000000022056366", 2, "i3"), ("000000000022069613", 4, "i4")],    # 22069613: fora da tabela
    WO3: [("000000000022056366", 1.5, "i5"), ("000000000022099999", 3, "i6")],  # 22099999: sem preço
}


def _transporte(volatil=None):
    leituras = {"n": 0}

    def t(url, timeout=None):
        wo = next(w for w in (WO1, WO2, WO3, WO4) if w in url.replace("%7C", "|"))
        leituras["n"] += 1
        itens = []
        for material, qtd, id_ in ITENS.get(wo, []):
            if id_ == volatil:
                qtd = qtd + leituras["n"]          # muda a cada leitura: divergente (não conferido)
            itens.append({"ID": id_, "workOrderID": wo, "material": material, "fornecedor": "0000137886",
                          "quantidadeFaltante": qtd, "mensagem": "Estoque indisponível"})
        return Resposta(200, "application/json", json.dumps({"value": itens}))
    return t


def _gerar(tmp_path, tabela, volatil=None):
    rs = processar_lote(ClienteBTP(_config(), _transporte(volatil)), [WO1, WO2, WO3, WO4], prefs=RAPIDO)
    destino = gerar_planilha(rs, tmp_path / "a.xlsx", tabela=tabela)
    assert verificar_planilha(destino, rs, tabela=tabela) == []
    return rs, destino, load_workbook(destino)


def test_coluna_tipo_aba_materiais_e_tabela_oculta(tmp_path):
    tabela = materiais.ler_planilha(tabela_xlsx(tmp_path / "t.xlsx", REAL), "TABELA VALORES 07-2026.xlsx")
    _, _, wb = _gerar(tmp_path, tabela)
    assert wb.sheetnames == ABAS and wb["Materiais"].sheet_state == "visible"
    assert wb["Tabela de valores"].sheet_state == "hidden" and wb["Rastreio"].sheet_state == "hidden"
    d = list(wb["Dados"].iter_rows(values_only=True))
    assert d[0][:3] == ("WO", "Material SAP", "Tipo Material")
    tipo = {(r[0], r[1]): r[2] for r in d[1:]}
    assert tipo[(WO1, "000000000022056366")] == "ISOLADOR COAXIAL QUADRADO - CISP-HR"
    assert tipo[(WO2, "000000000022069613")] == "-"                     # fora da tabela: não chuta
    assert tipo[(WO4, "-")] == "-"                                       # WO sem alerta

    m = list(wb["Materiais"].iter_rows(values_only=True))
    assert m[0] == ("Material SAP", "Tipo Material", "Quantidade total", "Valor total (R$)", "Valor unitário (R$)",
                    "Observação")
    linhas = {r[0]: r for r in m[1:] if r[0] and r[0] != "Total"}
    # 22056366 em 3 WOs: 1 + 2 + 1,5 = 4,5 un × R$ 15,43 = R$ 69,435
    assert linhas["000000000022056366"][1:5] == ("ISOLADOR COAXIAL QUADRADO - CISP-HR", 4.5, 69.435, 15.43)
    assert linhas["000000000022056366"][5] == "-"
    assert linhas["000000000022025349"][2:5] == (10, 1.3, 0.13)
    assert linhas["000000000022069613"][1:5] == ("-", 4, "-", "-") and "fora da tabela" in linhas["000000000022069613"][5]
    assert linhas["000000000022099999"][3:5] == ("-", "-") and "Sem valor unitário" in linhas["000000000022099999"][5]
    assert [r[0] for r in m[1:5]] == ["000000000022056366", "000000000022025349",   # maior valor primeiro
                                      "000000000022069613", "000000000022099999"]
    total = m[-1]
    assert total[0] == "Total" and total[3] == pytest.approx(70.735) and "2 material(is) sem valor" in total[5]
    assert wb["Materiais"]["D2"].number_format == wb["Materiais"]["E2"].number_format == '"R$" #,##0.00'

    t = list(wb["Tabela de valores"].iter_rows(values_only=True))
    assert t[0] == ("Material SAP", "Tipo Material", "Valor unitário (R$)")
    assert ("000000000022056366", "ISOLADOR COAXIAL QUADRADO - CISP-HR", 15.43) in t   # 18 dígitos: PROCV direto
    assert len(t) == 1 + len(REAL)
    assert "TABELA VALORES 07-2026.xlsx" in wb["Análise"]["A3"].value


def test_item_nao_conferido_fica_fora_da_soma(tmp_path):
    tabela = materiais.ler_planilha(tabela_xlsx(tmp_path / "t.xlsx", REAL))
    _, _, wb = _gerar(tmp_path, tabela, volatil="i3")       # WO2, 22056366: quantidade muda entre leituras
    m = {r[0]: r for r in wb["Materiais"].iter_rows(values_only=True) if r[0]}
    assert m["000000000022056366"][2] == 2.5                 # 1 + 1,5 (o divergente não entra)
    assert "1 item(ns) não conferido(s) fora da soma" in m["000000000022056366"][5]


def test_sem_tabela_tudo_sai_com_traco(tmp_path):
    _, _, wb = _gerar(tmp_path, None)
    assert {r[2] for r in wb["Dados"].iter_rows(min_row=2, values_only=True)} == {"-"}
    m = {r[0]: r for r in wb["Materiais"].iter_rows(values_only=True) if r[0]}
    assert m["000000000022056366"][1:5] == ("-", 4.5, "-", "-") and "não importada" in m["000000000022056366"][5]
    assert m["Total"][3] == "-"
    assert list(wb["Tabela de valores"].iter_rows(values_only=True)) == [("Material SAP", "Tipo Material",
                                                                          "Valor unitário (R$)")]


def test_verificacao_acusa_valor_alterado_na_aba_materiais(tmp_path):
    tabela = materiais.ler_planilha(tabela_xlsx(tmp_path / "t.xlsx", REAL))
    rs, destino, wb = _gerar(tmp_path, tabela)
    wb["Materiais"]["D2"] = 1
    wb.save(destino)
    assert any("Materiais" in p and "Valor total" in p for p in verificar_planilha(destino, rs, tabela=tabela))


# ============================================================ API
def test_api_importa_consulta_e_remove(tmp_path, monkeypatch):
    monkeypatch.setattr(cfg, "ARQUIVO_CONFIG", tmp_path / "config" / "api.json")
    from agente_btp.agente import Agente
    from agente_btp.servidor import criar_app
    c = criar_app(Agente()).test_client()
    assert c.get("/api/materiais/tabela").json == {"carregada": False}
    corpo = tabela_xlsx(tmp_path / "t.xlsx", REAL).read_bytes()
    r = c.post("/api/materiais/tabela", data={"arquivo": (io.BytesIO(corpo), "TABELA VALORES 07-2026.xlsx")},
               content_type="multipart/form-data")
    assert r.status_code == 200 and r.json["materiais"] == 5 and r.json["sem_valor"] == 1
    assert c.get("/api/materiais/tabela").json["arquivo"] == "TABELA VALORES 07-2026.xlsx"
    ruim = tabela_xlsx(tmp_path / "r.xlsx", [(1, "x")], cabecalho=("Material", "Descrição")).read_bytes()
    r = c.post("/api/materiais/tabela", data={"arquivo": (io.BytesIO(ruim), "r.xlsx")}, content_type="multipart/form-data")
    assert r.status_code == 400 and "Texto breve material" in r.json["erro"]
    assert c.get("/api/materiais/tabela").json["materiais"] == 5     # a tabela anterior continua valendo
    assert c.delete("/api/materiais/tabela").json["carregada"] is False
    assert c.get("/api/materiais/tabela").json == {"carregada": False}
