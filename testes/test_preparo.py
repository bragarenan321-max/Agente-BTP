"""Preparação automática: só aceita o caminho que traz o conteúdo da janela "Mensagens".

Cada cenário imita um jeito diferente de o BTP entregar as mensagens. Em todos, o caminho
"ingênuo" (filtrar o conjunto de mensagens pelo Consolidado) volta VAZIO — exatamente o
que aconteceu no BTP real (499 WOs como "Sem registros").
"""
import json
import time
from urllib.parse import unquote

import pytest

from agente_btp import agente as agente_mod
from agente_btp import config as cfg
from agente_btp import historico
from agente_btp.cliente import Resposta
from agente_btp.modelos import Status

WO = "06162|231715939"
CONSOL = "105012-20260916-mEeF+Ow/Dh"
ITEM = {"Wo": WO, "Matnr": "000000000022066463", "Lifnr": "0000138480", "QtdFaltante": "1.000",
        "Mensagem": "Estoque indisponível"}

V2 = """<?xml version="1.0" encoding="utf-8"?>
<edmx:Edmx Version="1.0" xmlns:edmx="http://schemas.microsoft.com/ado/2007/06/edmx"
 xmlns:m="http://schemas.microsoft.com/ado/2007/08/dataservices/metadata">
 <edmx:DataServices><Schema Namespace="ZMON" xmlns="http://schemas.microsoft.com/ado/2008/09/edm">
  <EntityType Name="Alerta"><Key><PropertyRef Name="ConsolidadoNum"/></Key>
   <Property Name="Wo" Type="Edm.String"/><Property Name="ConsolidadoNum" Type="Edm.String"/>
   <Property Name="Mensagem" Type="Edm.String"/>%(json)s%(nav)s</EntityType>
  %(msg_tipo)s
  %(assoc)s
  <EntityContainer Name="C" m:IsDefaultEntityContainer="true">
   <EntitySet Name="AlertaSet" EntityType="ZMON.Alerta"/>%(msg_set)s
   %(funcoes)s
  </EntityContainer></Schema></edmx:DataServices></edmx:Edmx>"""
MSG_TIPO = """<EntityType Name="Msg"><Key><PropertyRef Name="Matnr"/></Key>
   <Property Name="ConsolidadoNum" Type="Edm.String"/><Property Name="Matnr" Type="Edm.String"/>
   <Property Name="Lifnr" Type="Edm.String"/><Property Name="QtdFaltante" Type="Edm.Decimal"/>
   <Property Name="Mensagem" Type="Edm.String"/></EntityType>"""
MSG_SET = '<EntitySet Name="MsgSet" EntityType="ZMON.Msg"/>'
NAV = '<NavigationProperty Name="ToMsg" Relationship="ZMON.A_M" FromRole="A" ToRole="M"/>'
ASSOC = '<Association Name="A_M"><End Type="ZMON.Alerta" Role="A" Multiplicity="1"/><End Type="ZMON.Msg" Role="M" Multiplicity="*"/></Association>'


def metadata(json_=False, nav=False, msgs=True, funcoes=""):
    return V2 % {"json": '<Property Name="ItensJson" Type="Edm.String"/>' if json_ else "",
                 "nav": NAV if nav else "", "assoc": ASSOC if nav else "",
                 "msg_tipo": MSG_TIPO if msgs else "", "msg_set": MSG_SET if msgs else "",
                 "funcoes": funcoes}


def ok(corpo):
    return Resposta(200, "application/json", json.dumps(corpo))


class BTP:
    """BTP falso. `rotas`: função(caminho_decodificado) -> registros | None (None = 404)."""

    def __init__(self, xml, rotas, alerta_extra=None):
        self.xml, self.rotas, self.urls = xml, rotas, []
        self.alerta = dict({"Wo": WO, "ConsolidadoNum": CONSOL, "Mensagem": "Erro durante baixa"},
                           **(alerta_extra or {}))

    def __call__(self, url, accept=None, timeout=None):
        self.urls.append(url)
        caminho = unquote(url)
        if caminho.endswith("$metadata"):
            return Resposta(200, "application/xml", self.xml)
        if caminho.startswith("/srv/AlertaSet?"):
            return ok({"d": {"results": [self.alerta] if f"'{WO}'" in caminho else []}})
        registros = self.rotas(caminho)
        if registros is None:
            return Resposta(404, "application/json", '{"error":{"message":"nao existe"}}')
        return ok({"d": {"results": registros}})


@pytest.fixture
def agente(tmp_path, monkeypatch):
    monkeypatch.setattr(cfg, "ARQUIVO_CONFIG", tmp_path / "api.json")
    monkeypatch.setattr(historico, "ARQUIVO", tmp_path / "hist.json")
    monkeypatch.setattr(historico, "CONTADOR", tmp_path / "contador.txt")
    monkeypatch.setattr(agente_mod, "PASTA_SAIDAS", tmp_path / "saidas")
    a = agente_mod.Agente()
    a.ponte.fontes = lambda: [{"url": "/srv/"}]
    return a


def preparar(a, btp, wos=(WO,)):
    a.ponte.get = btp
    a._testados, a._achados, a._sondadas = set(), {}, set()
    return a._preparar_com_wos(list(wos))


def gravada():
    return json.loads(cfg.ARQUIVO_CONFIG.read_text(encoding="utf-8"))


def vazio_exceto(**rotas):
    def resp(caminho):
        for trecho, registros in rotas.items():
            if trecho in caminho:
                return registros
        return [] if caminho.startswith("/srv/MsgSet") else None
    return resp


# ---------------------------------------------------------------- o defeito real
def test_recusa_caminho_que_acha_a_wo_mas_nao_traz_mensagens(agente):
    btp = BTP(metadata(), vazio_exceto())
    assert preparar(agente, btp) is False
    assert not cfg.ARQUIVO_CONFIG.exists()          # antes: gravava e dava 499 "Sem registros"


# ---------------------------------------------------------------- caminhos novos
def test_link_de_navegacao_informado_pelo_servidor(agente):
    uri = "https://interno.claro:8443/sap/opu/odata/sap/ZMON/AlertaSet('X-1')/ToMsg"
    btp = BTP(metadata(nav=True), vazio_exceto(**{"/srv/AlertaSet('X-1')/ToMsg": [ITEM]}),
              {"ToMsg": {"__deferred": {"uri": uri}}})
    assert preparar(agente, btp) is True
    assert gravada()["campo_mensagens_link"] == "ToMsg"


def test_mensagens_em_json_dentro_da_linha(agente):
    itens = json.dumps([{"Material": "000000000022066463", "Fornecedor": "0000138480",
                         "QtdFalta": 1, "Mensagem": "Estoque indisponível"}])
    btp = BTP(metadata(json_=True, msgs=False), vazio_exceto(), {"ItensJson": itens})
    assert preparar(agente, btp) is True
    assert gravada()["campo_mensagens_json"] == "ItensJson"
    assert gravada()["origem"] == "sondagem"


def test_funcao_de_consulta_e_nunca_funcao_de_acao(agente):
    funcoes = """<FunctionImport Name="ReprocessarAlerta" ReturnType="Collection(ZMON.Msg)" EntitySet="MsgSet" m:HttpMethod="GET">
      <Parameter Name="ConsolidadoNum" Type="Edm.String" Mode="In"/></FunctionImport>
     <FunctionImport Name="EfetuarBaixa" ReturnType="Collection(ZMON.Msg)" EntitySet="MsgSet" m:HttpMethod="GET">
      <Parameter Name="ConsolidadoNum" Type="Edm.String" Mode="In"/></FunctionImport>
     <FunctionImport Name="GetMensagens" ReturnType="Collection(ZMON.Msg)" EntitySet="MsgSet" m:HttpMethod="GET">
      <Parameter Name="Consolidado" Type="Edm.String" Mode="In"/></FunctionImport>"""
    btp = BTP(metadata(funcoes=funcoes), vazio_exceto(**{f"/srv/GetMensagens?Consolidado='{CONSOL}'": [ITEM]}))
    assert preparar(agente, btp) is True
    assert "GetMensagens" in gravada()["url_mensagens"]
    assert not [u for u in btp.urls if "Reprocessar" in u or "Baixa" in u]


def test_mensagens_de_outra_wo_sao_recusadas(agente):
    outra = dict(ITEM, Wo="99999|999999999")
    btp = BTP(metadata(), vazio_exceto(**{"/srv/MsgSet?": [outra]}))
    assert preparar(agente, btp) is False


def test_chave_estrangeira_alerta_id_v4(agente):
    xml = """<?xml version="1.0"?><edmx:Edmx Version="4.0" xmlns:edmx="http://docs.oasis-open.org/odata/ns/edmx">
    <edmx:DataServices><Schema Namespace="S" xmlns="http://docs.oasis-open.org/odata/ns/edm">
     <EntityType Name="Alertas"><Key><PropertyRef Name="ID"/></Key><Property Name="ID" Type="Edm.Guid"/>
      <Property Name="Wo" Type="Edm.String"/><Property Name="Consolidado" Type="Edm.String"/></EntityType>
     <EntityType Name="Itens"><Key><PropertyRef Name="Item"/></Key><Property Name="Item" Type="Edm.String"/>
      <Property Name="alerta_ID" Type="Edm.Guid"/><Property Name="Matnr" Type="Edm.String"/>
      <Property Name="Lifnr" Type="Edm.String"/><Property Name="QtdFaltante" Type="Edm.Decimal"/>
      <Property Name="Mensagem" Type="Edm.String"/></EntityType>
     <EntityContainer Name="C"><EntitySet Name="Alertas" EntityType="S.Alertas"/>
      <EntitySet Name="Itens" EntityType="S.Itens"/></EntityContainer></Schema></edmx:DataServices></edmx:Edmx>"""
    guid = "11111111-2222-3333-4444-555555555555"

    def get(url, accept=None, timeout=None):
        c = unquote(url)
        if c.endswith("$metadata"):
            return Resposta(200, "application/xml", xml)
        if c.startswith("/srv/Alertas?"):
            return ok({"value": [{"ID": guid, "Wo": WO, "Consolidado": CONSOL}]})
        if c == f"/srv/Itens?$filter=alerta_ID eq {guid}":
            return ok({"value": [ITEM]})
        return ok({"value": []})
    assert preparar(agente, get) is True
    assert "alerta_ID" in gravada()["url_mensagens"]


def test_mensagens_em_outro_servico(agente):
    agente.ponte.fontes = lambda: [{"url": "/srv/"}, {"url": "/msg/"}]
    xml_msg = metadata().replace('Name="AlertaSet" EntityType="ZMON.Alerta"', 'Name="Nada" EntityType="ZMON.X"')

    def get(url, accept=None, timeout=None):
        c = unquote(url)
        if c == "/msg/$metadata":
            return Resposta(200, "application/xml", xml_msg)
        if c.endswith("$metadata"):
            return Resposta(200, "application/xml", metadata(msgs=False))
        if c.startswith("/srv/AlertaSet?"):
            return ok({"d": {"results": [{"Wo": WO, "ConsolidadoNum": CONSOL}]}})
        if c == f"/msg/MsgSet?$filter=ConsolidadoNum eq '{CONSOL}'":
            return ok({"d": {"results": [ITEM]}})
        return ok({"d": {"results": []}})
    assert preparar(agente, get) is True
    assert gravada()["url_mensagens"].startswith("/msg/MsgSet")


# ---------------------------------------------------------------- salvaguardas
def test_configuracao_de_versao_anterior_e_descartada(tmp_path):
    caminho = tmp_path / "api.json"
    caminho.write_text(json.dumps({"url_lista": "/s/A?$filter=W eq '{wo}'", "url_mensagens": "/s/M",
                                   "campos": {"material": "M", "fornecedor": "F", "quantidade_faltante": "Q",
                                              "mensagem": "T"}, "origem": "dicionario"}))
    assert cfg.descartar_se_antiga(caminho) is True and not caminho.exists()
    cfg.salvar({"url_lista": "/s/A?$filter=W eq '{wo}'", "url_mensagens": "/s/M",
                "campos": {"material": "M", "fornecedor": "F", "quantidade_faltante": "Q", "mensagem": "T"}},
               caminho)
    assert cfg.descartar_se_antiga(caminho) is False and caminho.exists()


def test_lote_quase_todo_sem_registros_refaz_a_configuracao(agente):
    wos = [WO, "06162|231715940", "06162|231715941"]
    cfg.salvar({"url_lista": "/srv/AlertaSet?$filter=Wo eq '{wo}'", "campo_consolidado_lista": "ConsolidadoNum",
                "url_mensagens": "/srv/MsgSet?$filter=ConsolidadoNum eq '{linha[ConsolidadoNum]}'",
                "campos": {"material": "Matnr", "fornecedor": "Lifnr", "quantidade_faltante": "QtdFaltante",
                           "mensagem": "Mensagem"}, "origem": "dicionario", "confirmada": True})

    def get(url, accept=None, timeout=None):
        c = unquote(url)
        if c.startswith("/srv/AlertaSet?"):
            wo = c.split("'")[1]
            return ok({"d": {"results": [{"Wo": wo, "ConsolidadoNum": CONSOL}]}})
        return ok({"d": {"results": []}})
    agente.ponte.get = get
    agente.ponte.agente_id, agente.ponte._aba_autenticada = "aba", True
    agente.ponte._visto_em = time.monotonic() + 3600
    agente._cmd_executar(wos, "teste")
    e = agente.estado()
    assert e["contagem"] == {Status.SEM_REGISTROS.value: 3}
    assert not cfg.ARQUIVO_CONFIG.exists()
    assert "Mensagens" in e["aviso"]
    assert e["ultimo_arquivo"] == "Análise 1 - 3 WOs.xlsx"
    assert historico.listar()[0]["numero"] == 1


from servico_real_falso import SRV, servico_real  # noqa: E402


def test_servico_real_acha_itens_pelos_dados(agente):
    get, urls = servico_real()
    agente.ponte.fontes = lambda: [{"recurso": SRV + "$metadata"}]
    assert preparar(agente, get) is True
    dados = gravada()
    assert dados["origem"] == "exploracao"
    assert "ErrosMaterial" in dados["url_lista"]
    assert dados["campos"]["material"] == "material" and dados["campos"]["mensagem"] == "mensagem"
    assert len(urls) < 120, len(urls)              # rápido: poucas consultas, não milhares
    assert agente._diag["dicionario"]["exploracao"]  # o diagnóstico registra o que foi visto

    from agente_btp.cliente import ClienteBTP
    from agente_btp.processador import processar_lote
    [r] = processar_lote(ClienteBTP(cfg.carregar(), get), [WO])
    assert r.status is Status.PROCESSADA
    l = r.linhas[0]
    assert (l.wo, l.material, l.fornecedor, l.quantidade_faltante, l.mensagem) == (
        WO, "000000000022066463", "0000138480", 2, "Estoque indisponível")   # faltante, não a total (3)


def test_campo_tipo_wo_nao_e_numero_da_wo():
    from agente_btp.campos import tipo_do_campo
    assert tipo_do_campo("tipoWo") is None and tipo_do_campo("statusWO") is None
    assert tipo_do_campo("workOrderID_workOrderID") == "wo" and tipo_do_campo("numWO") == "wo"


def test_servico_real_itens_ligados_pelo_consolidado(agente):
    get, urls = servico_real(ligado=True)
    agente.ponte.fontes = lambda: [{"recurso": SRV + "$metadata"}]
    assert preparar(agente, get) is True
    dados = gravada()
    assert "Pendencias" in dados["url_lista"] and "ErrosMaterial?$filter=consolidID" in dados["url_mensagens"]
    from agente_btp.cliente import ClienteBTP
    from agente_btp.processador import processar_lote
    [r] = processar_lote(ClienteBTP(cfg.carregar(), get), [WO])
    assert r.status is Status.PROCESSADA and r.linhas[0].material == "000000000022066463"
    assert r.linhas[0].consolidado == CONSOL


def test_exploracao_tenta_a_proxima_wo(agente):
    get, urls = servico_real()
    agente.ponte.fontes = lambda: [{"recurso": SRV + "$metadata"}]
    assert preparar(agente, get, wos=("11111|111111111", WO)) is True
    assert gravada()["origem"] == "exploracao"


# ---------------------------------------------------------------- conferência da fonte com o usuário
def _rodar_com_conferencia(agente, wos, respostas, prazo=20):
    """Roda a análise numa thread e responde à conferência da amostra, como o usuário na tela."""
    import threading
    t = threading.Thread(target=agente._cmd_executar, args=(wos, "teste"), daemon=True)
    t.start()
    amostras, fim = [], time.monotonic() + prazo
    while t.is_alive() and time.monotonic() < fim:
        e = agente.estado()
        if e["fase"] == "confirmando" and e["confirmacao"] and respostas:
            amostras.append(e["confirmacao"])
            agente.responder_fonte(respostas.pop(0))
        time.sleep(0.05)
    t.join(timeout=5)
    return amostras


def _conectar(agente, get):
    agente.ponte.get = get
    agente.ponte.agente_id, agente.ponte._aba_autenticada = "aba", True
    agente.ponte._visto_em = time.monotonic() + 3600


def test_isca_do_btp_real_nao_e_escolhida(agente):
    """Caso real de 15/10: tabela com material de 8 dígitos e 'estoque'/'bom' não pode vencer."""
    get, _ = servico_real()
    agente.ponte.fontes = lambda: [{"recurso": SRV + "$metadata"}]
    assert preparar(agente, get) is True
    assert "MateriaisLancadosStatus" not in gravada()["url_lista"]


def test_lote_so_comeca_depois_da_conferencia_e_recusa_troca_a_fonte(agente, monkeypatch, tmp_path):
    monkeypatch.setattr(agente_mod, "ARQUIVO_REJEITADAS", tmp_path / "recusadas.json")
    get, _ = servico_real()
    agente.ponte.fontes = lambda: [{"recurso": SRV + "$metadata"}]
    _conectar(agente, get)
    # Começa com a fonte ERRADA (a isca), como aconteceu no BTP real.
    cfg.salvar({"url_lista": SRV + "MateriaisLancadosStatus?$filter=workOrderID eq '{wo}'",
                "campo_consolidado_lista": "", "url_mensagens": None, "campo_mensagens_expandidas": ".",
                "campos": {"wo": "workOrderID", "material": "material", "fornecedor": "fornecedor",
                           "quantidade_faltante": "quantidade", "mensagem": "tipoErro"}, "origem": "exploracao"})
    amostras = _rodar_com_conferencia(agente, [WO], [{"resposta": "outra"}, {"resposta": "sim", "mensagem": "item"}])
    assert [i["mensagem_item"] for i in amostras[0]["itens"]] == ["estoque", "bom"]   # o usuário vê a isca
    assert amostras[1]["itens"][0]["material"] == "000000000022066463"                # e depois a certa
    assert gravada()["confirmada"] is True and "ErrosMaterial" in gravada()["url_lista"]
    e = agente.estado()
    assert e["fase"] == "concluido" and e["contagem"] == {Status.PROCESSADA.value: 1}
    from openpyxl import load_workbook
    d = list(load_workbook(agente_mod.PASTA_SAIDAS / e["ultimo_arquivo"])["Dados"].iter_rows(min_row=2, values_only=True))
    assert d == [(WO, "000000000022066463", "0000138480", 2, "Estoque indisponível", None, "Validada")]


def test_fonte_confirmada_nao_pergunta_de_novo(agente, monkeypatch, tmp_path):
    monkeypatch.setattr(agente_mod, "ARQUIVO_REJEITADAS", tmp_path / "recusadas.json")
    get, _ = servico_real()
    agente.ponte.fontes = lambda: [{"recurso": SRV + "$metadata"}]
    _conectar(agente, get)
    assert _rodar_com_conferencia(agente, [WO], [{"resposta": "sim"}])
    assert _rodar_com_conferencia(agente, [WO], [{"resposta": "outra"}], prazo=8) == []   # nenhuma pergunta
    assert agente.estado()["fase"] == "concluido"


def test_mensagem_do_alerta_pode_ser_escolhida(agente, monkeypatch, tmp_path):
    monkeypatch.setattr(agente_mod, "ARQUIVO_REJEITADAS", tmp_path / "recusadas.json")
    get, _ = servico_real(ligado=True)        # itens ligados ao alerta (Pendencias) pelo consolidado
    agente.ponte.fontes = lambda: [{"recurso": SRV + "$metadata"}]
    _conectar(agente, get)
    amostras = _rodar_com_conferencia(agente, [WO], [{"resposta": "sim", "mensagem": "alerta"}])
    assert amostras[0]["tem_mensagem_alerta"] and amostras[0]["itens"][0]["mensagem_alerta"]
    from openpyxl import load_workbook
    e = agente.estado()
    d = list(load_workbook(agente_mod.PASTA_SAIDAS / e["ultimo_arquivo"])["Dados"].iter_rows(min_row=2, values_only=True))
    assert d[0][4] == amostras[0]["itens"][0]["mensagem_alerta"]


def test_raio_x_mostra_de_onde_vem_cada_dado():
    from agente_btp.raiox import raio_x, resumo_texto
    get, urls = servico_real()

    class Aba:
        pagina = "/site#ObjMotor-display"

        def fontes(self):
            return [{"recurso": SRV + "$metadata"}]

        def get(self, url, accept=None):
            return get(url, accept)
    rel = raio_x(Aba(), WO, CONSOL)
    texto = resumo_texto(rel)
    assert "ErrosMaterial [workOrderID_workOrderID = wo]: 1 registro(s)" in texto
    assert "MateriaisLancadosStatus" in texto
    pend = next(c for c in rel["conjuntos"] if c["conjunto"] == "Pendencias")
    assert any(q["filtro"] == "consolidID = consolidado" for q in pend["consultas"])   # também pelo consolidado
    erros = next(c for c in rel["conjuntos"] if c["conjunto"] == "ErrosMaterial")
    assert erros["consultas"][0]["amostra"][0]["quantidadeFaltante"] == 2
    assert not [u for u in urls if "LimpaWorkOrder" in u]          # nome de ação nunca consultado


def test_caminho_da_janela_mensagens_confirmado_no_btp(agente):
    """Pendencias -> UltimoRetornoView com status 'E': só as linhas que a janela Mensagens mostra."""
    get, urls = servico_real(retorno=True)
    agente.ponte.fontes = lambda: [{"recurso": SRV + "$metadata"}]
    assert preparar(agente, get) is True
    dados = gravada()
    assert dados["origem"] == "receita_btp" and "status eq 'E'" in dados["url_mensagens"]
    from agente_btp.cliente import ClienteBTP
    from agente_btp.processador import processar_lote
    [r] = processar_lote(ClienteBTP(cfg.carregar(), get), [WO])
    assert sorted((l.material, l.quantidade_faltante, l.mensagem) for l in r.linhas) == [
        ("000000000022056366", 1, "Estoque indisponível"), ("000000000022060284", 1, "Estoque indisponível")]
    assert all(l.fornecedor == "0000148963" and l.consolidado == CONSOL for l in r.linhas)
    assert not any("22061736" in (l.material or "") for l in r.linhas)          # a linha "Ok" não entra


def test_observacao_da_tela_nao_passa_na_frente_do_caminho_confirmado(agente):
    """Caso real 27/10: a observação gravou RetornoDaBaixa (histórico, 2 quantidades) antes de o
    dicionário ser lido; todas as WOs ficaram incompletas. Com o caminho confirmado disponível,
    a observação não grava nada — e, sem o dicionário lido, também espera."""
    get, _ = servico_real(retorno=True)
    agente.ponte.get = get
    agente.ponte.fontes = lambda: [{"recurso": SRV + "$metadata"}]
    consultas = [{"metodo": "GET", "url": SRV + f"RetornoDaBaixa?$filter=workOrderID eq '{WO}'", "status": 200,
                  "content_type": "application/json", "corpo_resposta": json.dumps({"value": [
                      {"workOrderID": WO, "material": "000000000022056366", "fornecedorID": "0000148963",
                       "quantidade": "1e+0", "quantidadeDoc": "0e+0", "mensagem": "Estoque indisponível"}]})}]
    agente._ao_observar(consultas)                       # dicionário ainda não lido: espera
    assert not cfg.ARQUIVO_CONFIG.exists()
    agente._buscar_candidatos()
    assert agente._receita_disponivel()
    agente._ao_observar(consultas)                       # receita disponível: observação não decide
    assert not cfg.ARQUIVO_CONFIG.exists()


def test_fonte_nao_confirmada_e_trocada_pelo_caminho_confirmado_antes_da_amostra(agente):
    get, _ = servico_real(retorno=True)
    _conectar(agente, get)
    agente.ponte.fontes = lambda: [{"recurso": SRV + "$metadata"}]
    cfg.salvar({"url_lista": SRV + "MateriaisLancadosStatus?$filter=workOrderID eq '{wo}'",
                "campo_consolidado_lista": "", "url_mensagens": None, "campo_mensagens_expandidas": ".",
                "campos": {"wo": "workOrderID", "material": "material", "fornecedor": "fornecedor",
                           "quantidade_faltante": "quantidade", "mensagem": "tipoErro"}, "origem": "observacao"})
    amostras = _rodar_com_conferencia(agente, [WO], [{"resposta": "sim"}])
    assert "UltimoRetornoView" in amostras[0]["fonte"] and "status E" in amostras[0]["fonte"]
    assert gravada()["origem"] == "receita_btp" and gravada()["confirmada"] is True
    assert sorted(i["material"] for i in amostras[0]["itens"]) == ["000000000022056366", "000000000022060284"]
