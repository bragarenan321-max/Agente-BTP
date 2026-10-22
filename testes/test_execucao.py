"""Correção de raiz: preservação (R1), retomada (R2), sessão (R3), desempenho com equivalência (R4),
diagnóstico (R5) e coerência da tela (R6)."""
import json
import shutil
import subprocess
import sys
import threading
import time
from datetime import datetime
from collections import Counter
from pathlib import Path

import pytest
from openpyxl import load_workbook

from agente_btp import diagnostico
from agente_btp.cliente import CacheExecucao, ClienteBTP, Resposta, Ritmo, classificar
from agente_btp.exportacao import gerar_planilha, linhas_dados, verificar_planilha
from agente_btp.loja import Loja, desserializar, estado_da_wo, serializar
from agente_btp.modelos import FalhaComunicacao, SessaoExpirada
from agente_btp.motor import Motor
from agente_btp.preferencias import Preferencias, como_dict
from agente_btp.processador import processar_wo

from btp_lote_falso import BTPLote, LOGIN_HTML, config

RAIZ = Path(__file__).resolve().parent.parent


# ------------------------------------------------------------------ apoio
class Ambiente:
    """Loja + motor + BTP falso, sem navegador. `anunciar()` = a aba se anunciou de novo."""

    def __init__(self, btp, prefs, run_id="r1", loja=None, com_chave=False):
        self.btp, self.prefs, self.run_id = btp, prefs, run_id
        self.api = config(com_chave)
        self.loja = loja or Loja()
        self.anuncios = 0
        self.cache = CacheExecucao()
        if prefs.tamanho_lote <= 1:
            self.cache.lote_ativo = False
        self.ritmo = Ritmo(0, dormir=lambda s: None)
        if not self.loja.obter(run_id):
            self.loja.criar(run_id, 1, "teste", btp.wos, como_dict(prefs), {})
        self.motor = None

    def cliente(self):
        return ClienteBTP(self.api, self.btp, ritmo=self.ritmo, cache=self.cache,
                          ao_erro=lambda e: self.loja.problema(self.run_id, e.motivo, e.evidencia))

    def sonda(self):
        try:
            self.cliente()._get(self.api.url_lista.replace("{wo}", self.btp.wos[0]), "lista")
            return True, "", {}
        except SessaoExpirada as exc:
            return False, exc.motivo, exc.evidencia
        except FalhaComunicacao as exc:
            return None, exc.motivo, exc.evidencia

    def anunciar(self):
        self.anuncios += 1

    def rodar(self, em_segundo_plano=False):
        self.motor = Motor(self.loja, self.run_id, self.cliente, self.prefs, self.sonda,
                           anuncios=lambda: self.anuncios)
        if not em_segundo_plano:
            return self.motor.executar()
        self.resultado = None
        t = threading.Thread(target=lambda: setattr(self, "resultado", self.motor.executar()), daemon=True)
        t.start()
        return t


def esperar(condicao, limite=30.0):
    fim = time.monotonic() + limite
    while time.monotonic() < fim:
        if condicao():
            return True
        time.sleep(0.02)
    raise AssertionError("condição não atingida a tempo")


def campos(r):
    """O que a análise AFIRMA de uma WO, campo a campo."""
    return (r.wo, r.status.name, r.auditoria.name,
            [(l.wo, l.material, l.fornecedor, str(l.quantidade_faltante), l.mensagem, l.consolidado, l.chave,
              l.auditoria.name) for l in r.linhas])


SERIAL = Preferencias(concorrencia=1, tamanho_lote=1)


# ============================================================ R1 — preservação
def test_r1_queda_na_wo_51_preserva_as_50_validadas_e_retoma_sozinha():
    btp = BTPLote(100)
    btp.tipos = {w: "normal" for w in btp.wos}
    alvo = btp.wos[50]
    btp.cair_quando = lambda url: f"'{alvo.replace('|', '%7C')}'" in url and not getattr(btp, "_caiu", False) \
        and not setattr(btp, "_caiu", True)
    amb = Ambiente(btp, SERIAL)
    t = amb.rodar(em_segundo_plano=True)
    esperar(lambda: amb.loja.obter("r1")["situacao"] == "pausado_sessao")
    estados = amb.loja.estados_ordenados("r1")
    validadas_antes = {w: amb.loja.linha_wo("r1", w) for w in btp.wos[:50]}
    assert all(e in ("validada",) for e in estados[:50]), Counter(estados[:50])
    assert all(e == "pendente" for e in estados[50:])               # nenhuma virou falha
    assert amb.loja.contadores("r1")["falha"] == 0
    assert "falha" not in Counter(estados)
    # o usuário faz login de novo e a aba se anuncia: a análise continua sozinha
    btp.sessao_ok = True
    amb.anunciar()
    t.join(30)
    assert amb.resultado == "concluida"
    c = amb.loja.contadores("r1")
    assert c["validada"] + c["incompleta"] + c["divergente"] == 100 and c["falha"] == 0 and c["pendente"] == 0
    for w, antes in validadas_antes.items():                         # as 50 primeiras NÃO foram refeitas
        depois = amb.loja.linha_wo("r1", w)
        assert depois["hash_auditorias"] == antes["hash_auditorias"]
        assert depois["concluido_em"] == antes["concluido_em"] and depois["tentativas"] == 1
    assert amb.loja.linha_wo("r1", alvo)["tentativas"] == 2


@pytest.mark.parametrize("modo", ["redirecionar", "401", "403", "html"])
def test_r1_queda_no_meio_com_paralelismo_e_lote_nunca_vira_falha(modo):
    btp = BTPLote(300)
    btp.modo_queda = modo
    contagem = {"n": 0}

    def cair(url):
        contagem["n"] += 1
        return contagem["n"] == 40
    btp.cair_quando = cair
    amb = Ambiente(btp, Preferencias(concorrencia=4, tamanho_lote=20))
    t = amb.rodar(em_segundo_plano=True)
    esperar(lambda: amb.loja.obter("r1")["situacao"] == "pausado_sessao")
    esperar(lambda: amb.loja.contadores("r1")["em_andamento"] == 0)
    c = amb.loja.contadores("r1")
    assert c["falha"] == 0 and c["pendente"] > 0
    assert amb.loja.obter("r1")["motivo"] == {"redirecionar": "sessao_expirada", "401": "http_401",
                                             "403": "http_403", "html": "sessao_expirada"}[modo]
    btp.sessao_ok = True
    amb.anunciar()
    t.join(60)
    c = amb.loja.contadores("r1")
    assert amb.resultado == "concluida" and c["falha"] == 0 and c["pendente"] == 0 and c["em_andamento"] == 0


# ============================================================ R2 — retomada idempotente
def test_r2_retomar_so_refaz_o_que_falta_e_e_idempotente():
    btp = BTPLote(120)
    quebradas = btp.wos[5:8]
    for w in quebradas:
        btp.status_por_wo[w] = 503
    amb = Ambiente(btp, Preferencias(concorrencia=3, tamanho_lote=10, repeticoes_falha=0))
    assert amb.rodar() == "concluida"
    c = amb.loja.contadores("r1")
    assert c["falha"] == 3 and amb.loja.a_processar("r1") == quebradas
    assert {amb.loja.linha_wo("r1", w)["motivo"] for w in quebradas} == {"http_5xx"}
    fotos = {w: amb.loja.linha_wo("r1", w) for w in btp.wos if w not in quebradas}

    # 1ª retomada: o BTP voltou; só as 3 que falharam são consultadas de novo
    btp.status_por_wo.clear()
    btp.urls.clear()
    assert amb.rodar() == "concluida"
    consultadas = {w for w in btp.wos if any(f"'{w.replace('|', '%7C')}'" in u for u in btp.urls)}
    assert consultadas <= set(quebradas) | {btp.wos[0]}                 # + a WO da sonda
    assert amb.loja.contadores("r1")["falha"] == 0
    for w, antes in fotos.items():
        assert amb.loja.linha_wo("r1", w) == antes                     # nada reescrito

    # 2ª retomada: nada a fazer, nada muda (idempotente)
    tudo = {w: amb.loja.linha_wo("r1", w) for w in btp.wos}
    btp.urls.clear()
    assert amb.rodar() == "concluida"
    assert len(btp.urls) == 1                                          # só a sonda
    assert {w: amb.loja.linha_wo("r1", w) for w in btp.wos} == tudo


def test_r2_planilha_a_qualquer_momento_com_coluna_status(tmp_path):
    btp = BTPLote(40)
    amb = Ambiente(btp, SERIAL)
    btp.cair_quando = lambda url: btp.consultas == 30
    t = amb.rodar(em_segundo_plano=True)
    esperar(lambda: amb.loja.obter("r1")["situacao"] == "pausado_sessao")
    resultados = amb.loja.resultados("r1")
    destino = gerar_planilha(resultados, tmp_path / "parcial.xlsx", numero=4)
    assert verificar_planilha(destino, resultados) == []
    linhas = list(load_workbook(destino)["Dados"].iter_rows(values_only=True))
    assert linhas[0][-1] == "Status"
    por_wo = {}
    for l in linhas[1:]:
        por_wo[l[0]] = l[-1]
    c = amb.loja.contadores("r1")
    assert Counter(por_wo.values())["Pendente"] == c["pendente"] + c["em_andamento"] > 0
    amb.motor.parar()
    t.join(10)


def test_r2_api_de_retomada(tmp_path, monkeypatch):
    from agente_btp import agente as agente_mod, config as cfg, historico
    from agente_btp.servidor import criar_app
    monkeypatch.setattr(cfg, "ARQUIVO_CONFIG", tmp_path / "api.json")
    monkeypatch.setattr(historico, "ARQUIVO", tmp_path / "hist.json")
    monkeypatch.setattr(historico, "CONTADOR", tmp_path / "contador.txt")
    monkeypatch.setattr(agente_mod, "PASTA_SAIDAS", tmp_path / "saidas")
    btp = BTPLote(60)
    api = config()
    from dataclasses import asdict
    cfg.salvar(asdict(api))
    a = agente_mod.Agente()
    a.ponte.get = btp
    a.ponte.agente_id, a.ponte._aba_autenticada = "aba", True
    a.ponte._visto_em = time.monotonic() + 3600
    app = criar_app(a).test_client()
    btp.cair_quando = lambda url: btp.consultas == 100
    assert app.post("/api/executar", data={"texto": "\n".join(btp.wos)}).get_json()["ok"]
    esperar(lambda: a.estado().get("situacao") == "pausado_sessao")
    e = a.estado()
    assert e["conexao"] == "sessao_expirada"                         # honesto: sessão caiu
    assert e["execucao"]["contadores"]["falha"] == 0
    run_id = e["run_id"]
    a.parar()
    esperar(lambda: not a.estado()["ocupado"])
    ret = a.estado()["retomavel"]
    assert ret["run_id"] == run_id and ret["restantes"] > 0 and ret["validadas"] > 0
    btp.sessao_ok = True
    r = app.post("/api/executar", json={"run_id": run_id, "modo": "retomar"}).get_json()
    assert r["ok"] and r["a_fazer"] == ret["restantes"]
    esperar(lambda: not a.estado()["ocupado"] and a.estado()["fase"] == "concluido")
    c = a.loja.contadores(run_id)
    assert c["pendente"] == 0 and c["falha"] == 0
    assert a.estado()["retomavel"] is None
    # diagnóstico da análise (R5): JSON exportável, com a queda de sessão e a evidência
    d = app.get(f"/api/diagnostico?run_id={run_id}").get_json()
    sessao = next(p for p in d["problemas"] if p["motivo"] == "sessao_expirada")
    assert sessao["categoria"] == "FONTE" and sessao["primeira_ocorrencia"]["url"]
    assert sessao["primeira_ocorrencia"].get("redirecionado") is True
    assert any("PAUSADA" in t["evento"] for t in d["linha_do_tempo"])
    assert any("retomada" in t["evento"].lower() for t in d["linha_do_tempo"])
    assert historico.listar()[0]["run_id"] == run_id
    arq = app.get(f"/api/execucoes/{run_id}/planilha")
    assert arq.status_code == 200


# ============================================================ R3 — sessão
def test_r3_deteccao_de_sessao_morta():
    def erro(resp):
        with pytest.raises(Exception) as e:
            classificar(resp, "/srv/X")
        return type(e.value).__name__, e.value.motivo, e.value.evidencia
    assert erro(Resposta(401, "application/json", "{}"))[:2] == ("SessaoExpirada", "http_401")
    assert erro(Resposta(403, "text/plain", "Forbidden"))[:2] == ("SessaoExpirada", "http_403")
    assert erro(Resposta(0, "", "", redirecionado=True))[:2] == ("SessaoExpirada", "sessao_expirada")
    tipo, motivo, ev = erro(Resposta(200, "text/html", LOGIN_HTML))
    assert (tipo, motivo) == ("SessaoExpirada", "sessao_expirada") and ev["marcador_login"]
    assert "PHNhbWxwOkF1dGhuUmVx" not in json.dumps(ev)                # SAMLRequest não vaza
    # HTML de login com content-type "errado" (json) também é sessão
    assert erro(Resposta(200, "application/json", LOGIN_HTML))[1] == "sessao_expirada"
    assert erro(Resposta(429, "text/plain", "", retry_after="3"))[:2] == ("FalhaComunicacao", "http_429")
    assert erro(Resposta(503, "text/plain", ""))[:2] == ("FalhaComunicacao", "http_5xx")
    assert classificar(Resposta(200, "application/json", '{"value":[]}')) == '{"value":[]}'
    # "login" numa mensagem do BTP (JSON) NÃO é sessão
    assert classificar(Resposta(200, "application/json", '{"value":[{"mensagem":"erro no login do técnico"}]}'))


def test_r3_403_so_numa_wo_com_sessao_valida_nao_pausa():
    btp = BTPLote(30)
    btp.status_por_wo[btp.wos[4]] = 403
    amb = Ambiente(btp, Preferencias(concorrencia=2, tamanho_lote=5))
    assert amb.rodar() == "concluida"
    linha = amb.loja.linha_wo("r1", btp.wos[4])
    assert linha["status"] == "falha" and linha["motivo"] == "http_403"
    c = amb.loja.contadores("r1")
    assert c["falha"] == 1 and c["pendente"] == 0


def test_r3_sonda_antes_de_comecar_pausa_sem_consultar_wos():
    btp = BTPLote(10)
    btp.sessao_ok = False
    amb = Ambiente(btp, SERIAL)
    t = amb.rodar(em_segundo_plano=True)
    esperar(lambda: amb.loja.obter("r1")["situacao"] == "pausado_sessao")
    assert amb.loja.contadores("r1")["pendente"] == 10
    assert len(btp.urls) == 1                                          # só a sonda
    btp.sessao_ok = True
    amb.anunciar()
    t.join(20)
    assert amb.resultado == "concluida"


def test_r3_conexao_honesta(tmp_path, monkeypatch):
    from agente_btp import agente as agente_mod
    a = agente_mod.Agente()
    assert a.estado_conexao()["conexao"] == "desconectado"            # nunca falou
    a.ponte.agente_id, a.ponte._visto_em = "aba", time.monotonic()
    assert a.estado_conexao()["conexao"] == "aba_conectada"           # batimento sim, sessão não verificada
    a._sessao_ok()
    assert a.estado_conexao()["conexao"] == "conectado"
    a._sessao["em"] = time.monotonic() - agente_mod.VALIDADE_SESSAO_SEG - 1
    assert a.estado_conexao()["conexao"] == "aba_conectada"           # verificação velha não vale
    a._sessao_caiu("http_401")
    assert a.estado_conexao()["conexao"] == "sessao_expirada"
    a.ponte._visto_em = time.monotonic() - 999                        # sem batimento recente
    assert a.estado_conexao()["conexao"] == "desconectado"


def test_r3_ritmo_adaptativo_respeita_429():
    pausas = []
    ritmo = Ritmo(0, dormir=pausas.append)
    ritmo.depois(429, "2")
    ritmo.antes()
    ritmo.antes()
    assert ritmo.espera >= 2 and pausas and max(pausas) >= 1.9
    for _ in range(60):
        ritmo.depois(200)
    assert ritmo.espera == 0                                          # volta ao normal sem pausa fixa


def test_r3_429_no_lote_nao_vira_falha():
    btp = BTPLote(40)
    btp.erros_429 = 3
    amb = Ambiente(btp, Preferencias(concorrencia=2, tamanho_lote=10))
    assert amb.rodar() == "concluida"
    assert amb.loja.contadores("r1")["falha"] == 0
    assert any(p["motivo"] == "http_429" for p in amb.loja.problemas("r1"))


# ============================================================ R4 — desempenho e equivalência
def _versao_antiga(destino: Path):
    """Extrai o pacote da versão anterior (commit 72c9acd) como agente_btp_antigo."""
    if not (RAIZ / ".git").exists() or not shutil.which("git"):
        pytest.skip("sem o histórico git para comparar com a versão anterior")
    pasta = destino / "agente_btp_antigo"
    pasta.mkdir()
    lista = subprocess.run(["git", "-C", str(RAIZ), "ls-tree", "--name-only", "72c9acd", "agente_btp/"],
                           capture_output=True, text=True)
    if lista.returncode:
        pytest.skip("commit da versão anterior indisponível")
    for nome in lista.stdout.split():
        if nome.endswith(".py"):
            conteudo = subprocess.run(["git", "-C", str(RAIZ), "show", f"72c9acd:{nome}"], capture_output=True,
                                      text=True, check=True).stdout
            (pasta / Path(nome).name).write_text(conteudo, encoding="utf-8")
    sys.path.insert(0, str(destino))
    import importlib
    return (importlib.import_module("agente_btp_antigo.processador"),
            importlib.import_module("agente_btp_antigo.cliente"),
            importlib.import_module("agente_btp_antigo.config"),
            importlib.import_module("agente_btp_antigo.preferencias"),
            importlib.import_module("agente_btp_antigo.exportacao"))


def test_r4_equivalencia_campo_a_campo_em_500_wos_contra_a_versao_anterior(tmp_path, monkeypatch):
    proc_v, cli_v, cfg_v, pref_v, exp_v = _versao_antiga(tmp_path)
    monkeypatch.setattr(proc_v, "_dormir", lambda s: None)
    btp_velho, btp_novo = BTPLote(500, pagina=50), BTPLote(500, pagina=50)
    from dataclasses import asdict
    api_velha = cfg_v.a_partir_de_dict(asdict(config()))
    antigos = proc_v.processar_lote(cli_v.ClienteBTP(api_velha, btp_velho), btp_velho.wos,
                                    prefs=pref_v.Preferencias())             # padrão antigo: 5 auditorias
    amb = Ambiente(btp_novo, Preferencias())                                 # padrão novo: 2 + lote + 4 em paralelo
    assert amb.rodar() == "concluida"
    novos = amb.loja.resultados("r1")
    assert len(antigos) == len(novos) == 500
    diferentes = [(a.wo, campos(a), campos(n)) for a, n in zip(antigos, novos)
                  if (a.wo, a.status.name, a.auditoria.name,
                      [(l.wo, l.material, l.fornecedor, str(l.quantidade_faltante), l.mensagem, l.consolidado,
                        l.chave, l.auditoria.name) for l in a.linhas]) != campos(n)]
    assert diferentes == []
    # a aba Dados (o que o usuário recebe) é idêntica nas 6 colunas de dados
    from agente_btp.exportacao import _txt
    velhas = [[_txt(x) for x in l] for l in exp_v.linhas_dados(antigos)]
    assert [[_txt(x) for x in l[:6]] for l in linhas_dados(novos)] == velhas
    assert amb.cache.lote_ativo, amb.cache.lote_motivo                       # o lote foi validado no caminho
    # menos consultas ao BTP: relatório antes/depois
    print(f"\nconsultas: versão anterior {btp_velho.consultas} | nova {btp_novo.consultas}")
    assert btp_novo.consultas * 5 < btp_velho.consultas


def test_r4_equivalencia_lote_paralelo_x_individual_serial_com_releitura_por_chave():
    a, b = BTPLote(500, semente=11, pagina=30), BTPLote(500, semente=11, pagina=30)
    serial = Ambiente(a, Preferencias(concorrencia=1, tamanho_lote=1, auditorias_por_wo=3), com_chave=True)
    rapido = Ambiente(b, Preferencias(concorrencia=4, tamanho_lote=20, auditorias_por_wo=3), com_chave=True)
    assert serial.rodar() == "concluida" and rapido.rodar() == "concluida"
    assert [campos(r) for r in serial.loja.resultados("r1")] == [campos(r) for r in rapido.loja.resultados("r1")]
    assert [l.tipo for l in rapido.loja.resultados("r1")[4].leituras][:4] == [
        "Extração inicial", "Nova consulta", "Releitura por chave", "Nova consulta"]


def test_r4_lote_que_diverge_da_individual_e_desativado_e_resultado_segue_exato():
    """Servidor que ignora parte do filtro em lote (só responde a 1ª WO): a conferência pega."""
    btp = BTPLote(60)
    original = btp.__call__

    def so_primeira(url, accept=None, timeout=None):
        if "%20or%20" in url and "Pendencias" in url:
            return original(url.split("%20or%20")[0] + ")", accept, timeout)
        return original(url, accept, timeout)
    amb = Ambiente(btp, Preferencias(concorrencia=2, tamanho_lote=10))
    amb.btp = so_primeira
    amb.btp.wos = btp.wos
    assert amb.rodar() == "concluida"
    assert not amb.cache.lote_ativo
    ref = Ambiente(BTPLote(60), SERIAL)
    ref.rodar()
    assert [campos(r) for r in amb.loja.resultados("r1")] == [campos(r) for r in ref.loja.resultados("r1")]
    assert any(p["motivo"] == "lote_desativado" for p in amb.loja.problemas("r1"))


def test_r4_divergencia_escalona_para_5_leituras_e_nunca_valida():
    btp = BTPLote(30)
    vol = btp.wos[4]
    btp.volateis = {vol}
    amb = Ambiente(btp, Preferencias(concorrencia=2, tamanho_lote=10))
    assert amb.rodar() == "concluida"
    r = next(x for x in amb.loja.resultados("r1") if x.wo == vol)
    assert r.auditoria.name == "DIVERGENTE"
    tipos = [l.tipo for l in r.leituras]
    assert tipos.count("Leitura extra") == 3 and len(tipos) == 6         # 1 + 2 auditorias + 3 extras
    assert amb.loja.linha_wo("r1", vol)["motivo"] == "divergencia_persistente"
    assert "leituras extras individuais" in amb.loja.linha_wo("r1", vol)["origem_dado"]


# ============================================================ R5 — diagnóstico
def test_r5_excecao_do_programa_vira_motivo_com_traceback(monkeypatch):
    import agente_btp.processador as proc
    btp = BTPLote(12)
    original = proc._concluir

    def quebra(cliente, wo, *a, **k):
        if wo == btp.wos[4]:
            raise KeyError("campo_que_nao_existe")
        return original(cliente, wo, *a, **k)
    monkeypatch.setattr(proc, "_concluir", quebra)
    amb = Ambiente(btp, Preferencias(concorrencia=1, tamanho_lote=4))
    assert amb.rodar() == "concluida"
    [p] = [p for p in amb.loja.problemas("r1") if p["motivo"] == "excecao_app"]
    ev = p["primeira"]
    assert "KeyError" in ev["erro"] and "Traceback" in ev["traceback"]
    assert ev["origem"]["modulo"].startswith("agente_btp.") and ev["origem"]["linha"]
    assert amb.loja.linha_wo("r1", btp.wos[4])["motivo"] == "excecao_app"


def test_r5_taxonomia_e_sanitizacao():
    for motivo in ("sessao_expirada", "http_401", "http_403", "http_429", "http_5xx", "timeout_wo", "timeout_lote",
                   "wo_inexistente", "campo_ausente:material", "parse_falhou", "divergencia_persistente",
                   "excecao_app"):
        assert diagnostico.categoria(motivo) in ("FONTE", "APP", "ENTRADA")
        assert diagnostico.proxima_acao(motivo)
    assert diagnostico.categoria("wo_inexistente") == "ENTRADA" and diagnostico.categoria("excecao_app") == "APP"
    texto = diagnostico.sanitizar('Authorization: Bearer abc.def.ghi; access_token=SEGREDO; '
                                  'eyJhbGciOiJIUzI1NiJ9.eyJzdWIiOiIxMjM0In0.abcdefghijk')
    assert "SEGREDO" not in texto and "abc.def" not in texto and "eyJhbGci" not in texto


# ============================================================ R6 — coerência tela x base x planilha
def test_r6_contadores_registro_e_planilha_concordam(tmp_path):
    btp = BTPLote(150)
    btp.status_por_wo[btp.wos[7]] = 503
    btp.volateis = {btp.wos[14]}
    contagem = {"n": 0}

    def cair(url):
        contagem["n"] += 1
        return contagem["n"] == 60
    btp.cair_quando = cair
    amb = Ambiente(btp, Preferencias(concorrencia=3, tamanho_lote=8, repeticoes_falha=0))
    t = amb.rodar(em_segundo_plano=True)
    esperar(lambda: amb.loja.obter("r1")["situacao"] == "pausado_sessao")
    btp.sessao_ok = True
    amb.anunciar()
    t.join(60)
    c = amb.loja.contadores("r1")
    # registro: o ÚLTIMO evento de cada WO tem o mesmo estado da base
    ultimo = {}
    for e in amb.loja.eventos("r1", tipos=("wo",), limite=10000):
        ultimo[e["wo"]] = e["status"]
    assert ultimo == {w: s for w, s in amb.loja.estados("r1").items() if s not in ("pendente", "em_andamento")}
    assert Counter(ultimo.values()) == Counter({k: v for k, v in c.items() if v and k not in ("pendente", "em_andamento")})
    # planilha: a coluna Status de cada WO bate com a base
    resultados = amb.loja.resultados("r1")
    destino = gerar_planilha(resultados, tmp_path / "x.xlsx")
    status = {}
    for l in list(load_workbook(destino)["Dados"].iter_rows(values_only=True))[1:]:
        status[l[0]] = l[-1].lower()
    assert status == amb.loja.estados("r1")
    assert [estado_da_wo(r) for r in resultados] == amb.loja.estados_ordenados("r1")


def test_r6_estado_da_tela_vem_da_base(tmp_path, monkeypatch):
    from agente_btp import agente as agente_mod
    a = agente_mod.Agente()
    amb = Ambiente(BTPLote(50), Preferencias(concorrencia=2, tamanho_lote=10), loja=a.loja)
    amb.rodar()
    a._estado.update(run_id="r1", ultimo_numero=1, fase="concluido")
    e = a.estado()
    c = a.loja.contadores("r1")
    assert e["execucao"]["contadores"] == c
    assert e["progresso"] == {"atual": e["execucao"]["feitas"], "total": 50}
    assert e["indicadores"]["validadas"] == c["validada"] and e["indicadores"]["falhas"] == c["falha"]
    linhas_wo = [m for m in e["mensagens"] if "] " in m and "auditoria:" in m]
    assert len(linhas_wo) == 50


def test_serializacao_ida_e_volta_mantem_a_planilha(tmp_path):
    btp = BTPLote(80)
    btp.volateis = {btp.wos[4]}
    cli = ClienteBTP(config(), btp)
    originais = [processar_wo(cli, w, Preferencias()) for w in btp.wos[:30]]
    voltas = [desserializar(json.loads(json.dumps(serializar(r)))) for r in originais]
    assert [campos(r) for r in voltas] == [campos(r) for r in originais]
    assert linhas_dados(voltas) == linhas_dados(originais)
    assert [r.rodadas_concluidas for r in voltas] == [r.rodadas_concluidas for r in originais]
    assert [r.registros_auditoria for r in voltas] == [r.registros_auditoria for r in originais]


# ============================================================ ponte e tela
def test_ponte_com_varios_trabalhadores_na_aba_entrega_cada_resposta_a_quem_pediu():
    from agente_btp.ponte import PonteAba
    ponte = PonteAba()
    ponte.registrar_agente("aba", "/site")
    parar = threading.Event()

    def trabalhador():
        resposta = None
        while not parar.is_set():
            r = ponte.ciclo("aba", True, "/site", resposta=resposta, espera=0.2)
            resposta = None
            if r.get("tarefa"):
                t = r["tarefa"]
                time.sleep(0.01)
                resposta = {"id": t["id"], "status": 200, "ct": "application/json", "body": t["url"], "ra": "1"}
    abas = [threading.Thread(target=trabalhador, daemon=True) for _ in range(4)]
    for t in abas:
        t.start()
    saidas = {}

    def pedir(i):
        saidas[i] = ponte.get(f"/srv/X?i={i}")
    pedidos = [threading.Thread(target=pedir, args=(i,)) for i in range(40)]
    inicio = time.monotonic()
    for t in pedidos:
        t.start()
    for t in pedidos:
        t.join(10)
    parar.set()
    assert all(saidas[i].corpo == f"/srv/X?i={i}" for i in range(40))
    assert saidas[0].retry_after == "1"
    assert time.monotonic() - inicio < 40 * 0.01 * 0.8                 # houve paralelismo de fato
    assert ponte.anuncios == 1


def test_retomada_nao_oferece_wo_de_numero_invalido_e_apagar_historico_remove_a_analise(tmp_path, monkeypatch):
    from agente_btp import agente as agente_mod, historico
    from agente_btp.servidor import criar_app
    monkeypatch.setattr(historico, "ARQUIVO", tmp_path / "hist.json")
    monkeypatch.setattr(historico, "CONTADOR", tmp_path / "contador.txt")
    a = agente_mod.Agente()
    btp = BTPLote(20)
    btp.wos = btp.wos + ["123"]
    amb = Ambiente(btp, SERIAL, loja=a.loja)
    assert amb.rodar() == "concluida"
    assert amb.loja.linha_wo("r1", "123")["motivo"] == "wo_invalida"
    assert a.loja.a_processar("r1") == [] and a.estado()["retomavel"] is None
    historico.registrar({"id": "r1", "run_id": "r1", "numero": 1})
    app = criar_app(a).test_client()
    assert app.delete("/api/historico/r1").get_json()["apagadas"] == 1
    assert a.loja.obter("r1") is None


def test_ola_informa_o_paralelismo_configurado(tmp_path, monkeypatch):
    from agente_btp import agente as agente_mod, preferencias
    from agente_btp.config import URL_LAUNCHPAD
    from agente_btp.servidor import criar_app
    from urllib.parse import urlsplit
    monkeypatch.setattr(preferencias, "ARQUIVO", tmp_path / "p.json")
    preferencias.salvar({"concorrencia": 8})
    app = criar_app(agente_mod.Agente()).test_client()
    u = urlsplit(URL_LAUNCHPAD)
    r = app.post("/api/agente/ola", json={"agente": "x", "pagina": "/"}, headers={"Origin": f"{u.scheme}://{u.netloc}"})
    assert r.get_json()["trabalhadores"] == 6                          # limite de conexões do navegador


def test_preferencias_da_versao_1_migram_para_a_politica_nova(tmp_path):
    from agente_btp import preferencias
    arq = tmp_path / "p.json"
    arq.write_text(json.dumps({"auditorias_por_wo": 5, "consultas_extras_divergencia": 2,
                               "intervalo_auditorias_seg": 0.5, "tempo_max_wo_seg": 200}))
    p = preferencias.carregar(arq)
    assert (p.auditorias_por_wo, p.consultas_extras_divergencia, p.intervalo_auditorias_seg) == (2, 3, 0.0)
    assert p.tempo_max_wo_seg == 200                                   # o resto é mantido
    arq.write_text(json.dumps({"auditorias_por_wo": 7, "versao": 1}))
    assert preferencias.carregar(arq).auditorias_por_wo == 7           # escolha explícita do usuário fica


def test_excel_de_lote_grande_nao_e_quadratico(tmp_path, monkeypatch):
    """ws.max_row percorre a planilha inteira: chamá-lo por linha tornava o Excel de 13 mil WOs
    (≈40 mil linhas) impossível de gerar. O número de chamadas não pode crescer com as linhas."""
    from openpyxl.worksheet.worksheet import Worksheet
    from agente_btp.modelos import Auditoria, LinhaMensagem, ResultadoWO, Status
    chamadas = {"n": 0}
    original = Worksheet.max_row

    def contar(self):
        chamadas["n"] += 1
        return original.fget(self)
    monkeypatch.setattr(Worksheet, "max_row", property(contar))
    resultados = [ResultadoWO(f"{i:05d}|{i:09d}", Status.PROCESSADA, auditoria=Auditoria.VALIDADO, linhas=[
        LinhaMensagem("c", f"{i:05d}|{i:09d}", f"{i:018d}", "0000000001", 1, "Estoque indisponível",
                      auditoria=Auditoria.VALIDADO)]) for i in range(3000)]
    gerar_planilha(resultados, tmp_path / "grande.xlsx")
    assert chamadas["n"] < 30, chamadas["n"]


def test_resultados_leves_geram_a_mesma_planilha_e_historico(tmp_path):
    btp = BTPLote(60)
    btp.volateis = {btp.wos[3]}
    amb = Ambiente(btp, Preferencias())
    amb.rodar()
    completos, leves = amb.loja.resultados("r1"), amb.loja.resultados("r1", leve=True)
    assert linhas_dados(leves) == linhas_dados(completos)
    assert [r.rodadas_concluidas for r in leves] == [r.rodadas_concluidas for r in completos]
    assert [r.registros_auditoria for r in leves] == [r.registros_auditoria for r in completos]
    a = gerar_planilha(completos, tmp_path / "a.xlsx", emitida_em=datetime(2026, 1, 1))
    b = gerar_planilha(leves, tmp_path / "b.xlsx", emitida_em=datetime(2026, 1, 1))
    for aba in ("Dados", "Análise", "Rastreio"):
        assert list(load_workbook(a)[aba].iter_rows(values_only=True)) == \
            list(load_workbook(b)[aba].iter_rows(values_only=True))
